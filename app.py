from __future__ import annotations

import json
import math
import os
import re
import tkinter as tk
from dataclasses import dataclass
from pathlib import Path
from tkinter import colorchooser, filedialog, messagebox, simpledialog, ttk

from PIL import Image, ImageChops, ImageDraw, ImageFont, ImageOps, ImageTk
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from translations import LANGUAGES, translate


PALETTE = (
    "#ff4040",
    "#00b7ff",
    "#39c96b",
    "#ffb000",
    "#b45cff",
    "#ff65b5",
    "#00c9ad",
    "#f07828",
)
DPI_CHOICES = ("100", "200", "300", "400", "500", "600")
LASSO_PREVIEW_OPACITY = 15
VERTEX_MAGNET_RADIUS = 12
PROJECT_FORMAT = "image-analyzer-project"
PROJECT_VERSION = 1


@dataclass
class Region:
    name: str
    points: list[tuple[float, float]]
    color: str
    opacity: int = 40
    border_width: int = 2
    visible: bool = True
    percentage_base_index: int | None = None
    percentage_color: str | None = None
    percentage_label_position: tuple[float, float] | None = None


def calculate_region_areas(
    image_size: tuple[int, int], regions: list[Region]
) -> tuple[list[int], int]:
    """Return raster pixel areas per region and for the union of all regions."""
    union = Image.new("L", image_size, 0)
    region_areas: list[int] = []

    for region in regions:
        mask = Image.new("L", image_size, 0)
        if len(region.points) >= 3:
            ImageDraw.Draw(mask).polygon(region.points, fill=255)
        region_areas.append(mask.histogram()[255])
        union = ImageChops.lighter(union, mask)

    union_area = union.histogram()[255]
    return region_areas, union_area


def calculate_area_percentage(area: int, base_area: int) -> float | None:
    if base_area <= 0:
        return None
    return area / base_area * 100


def _percentage_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    try:
        return ImageFont.truetype("arial.ttf", size)
    except OSError:
        return ImageFont.truetype("DejaVuSans.ttf", size)


def _default_percentage_label_position(
    image_size: tuple[int, int], region: Region, font: ImageFont.ImageFont, text: str
) -> tuple[float, float]:
    mask = Image.new("L", image_size, 0)
    if len(region.points) >= 3:
        ImageDraw.Draw(mask).polygon(region.points, fill=255)
    bounds = mask.getbbox()
    if bounds is None:
        return 0.0, 0.0
    left, top, right, bottom = bounds
    text_bounds = ImageDraw.Draw(Image.new("L", (1, 1))).textbbox(
        (0, 0), text, font=font
    )
    text_width = text_bounds[2] - text_bounds[0]
    text_height = text_bounds[3] - text_bounds[1]
    padding = 3
    max_x = max(left, right - text_width - padding)
    max_y = max(top, bottom - text_height - padding)
    return float(min(left + padding, max_x)), float(min(top + padding, max_y))


def _percentage_labels(
    image_size: tuple[int, int], regions: list[Region]
) -> list[str | None]:
    areas, _ = calculate_region_areas(image_size, regions)
    labels: list[str | None] = []
    for region, area in zip(regions, areas):
        base_index = region.percentage_base_index
        base_area = (
            image_size[0] * image_size[1]
            if base_index is None
            else areas[base_index]
        )
        percentage = calculate_area_percentage(area, base_area)
        labels.append("—" if percentage is None else f"{percentage:.2f}%")
    return labels


def _draw_percentage_labels(
    image: Image.Image,
    regions: list[Region],
    selected_indices: list[int],
    labels: list[str | None],
    font_size: int,
    label_positions: dict[int, tuple[float, float]] | None = None,
    text_colors: dict[int, str] | None = None,
) -> list[tuple[int, tuple[float, float], tuple[int, int, int, int], str]]:
    draw = ImageDraw.Draw(image)
    font = _percentage_font(font_size)
    layout = []
    for index in selected_indices:
        if not regions[index].visible or labels[index] is None:
            continue
        region = regions[index]
        text = labels[index]
        position = (
            (label_positions or {}).get(index)
            or region.percentage_label_position
            or _default_percentage_label_position(image.size, region, font, text)
        )
        bounds = draw.textbbox(position, text, font=font)
        color = (text_colors or {}).get(index) or region.percentage_color or region.color
        red, green, blue = (
            int(color[offset : offset + 2], 16) for offset in (1, 3, 5)
        )
        stroke = "#000000" if red + green + blue > 380 else "#ffffff"
        draw.text(
            position,
            text,
            font=font,
            fill=color,
            stroke_width=max(1, font_size // 12),
            stroke_fill=stroke,
        )
        layout.append((index, position, bounds, text))
    return layout


def _region_export_crop_bounds(
    image_size: tuple[int, int],
    regions: list[Region],
    selected_indices: list[int],
    crop_to_regions: bool,
) -> tuple[int, int, int, int] | None:
    width, height = image_size
    coverage = Image.new("L", image_size, 0)
    selected_regions = [
        regions[index] for index in selected_indices if regions[index].visible
    ]
    for region in selected_regions:
        mask = Image.new("L", image_size, 0)
        if len(region.points) >= 3:
            ImageDraw.Draw(mask).polygon(region.points, fill=1)
        coverage = ImageChops.add(coverage, mask)
    if not crop_to_regions:
        return (0, 0, width, height)
    bounds = coverage.point(lambda value: 255 if value else 0).getbbox()
    if bounds is None:
        return None
    margin = max((region.border_width + 1) // 2 for region in selected_regions)
    left, top, right, bottom = bounds
    return (
        max(0, left - margin),
        max(0, top - margin),
        min(width, right + margin),
        min(height, bottom + margin),
    )


def make_annotated_image(image: Image.Image, regions: list[Region]) -> Image.Image:
    """Composite translucent region fills and opaque outlines over the source."""
    base = image.convert("RGBA")
    overlay = Image.new("RGBA", base.size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay, "RGBA")

    for region in regions:
        if not region.visible or len(region.points) < 3:
            continue
        red, green, blue = tuple(
            int(region.color[index : index + 2], 16) for index in (1, 3, 5)
        )
        fill = (red, green, blue, round(255 * region.opacity / 100))
        draw.polygon(
            region.points,
            fill=fill,
            outline=(red, green, blue, 255),
            width=max(1, region.border_width),
        )

    return Image.alpha_composite(base, overlay).convert("RGB")


def make_region_export(
    image: Image.Image,
    regions: list[Region],
    selected_indices: list[int],
    include_intersections: bool = True,
    crop_to_regions: bool = True,
    show_percentages: bool = False,
    percentage_font_size: int = 24,
    percentage_labels: list[str | None] | None = None,
    label_positions: dict[int, tuple[float, float]] | None = None,
    text_colors: dict[int, str] | None = None,
) -> Image.Image:
    """Render selected regions on the source image, optionally cropping the result."""
    selected_regions = [
        regions[index] for index in selected_indices if regions[index].visible
    ]
    result = make_annotated_image(image, selected_regions)
    coverage = Image.new("L", image.size, 0)
    for region in selected_regions:
        mask = Image.new("L", image.size, 0)
        if len(region.points) >= 3:
            ImageDraw.Draw(mask).polygon(region.points, fill=1)
        coverage = ImageChops.add(coverage, mask)

    if not include_intersections:
        overlap_mask = coverage.point(lambda value: 255 if value > 1 else 0)
        result.paste(image.convert("RGB"), mask=overlap_mask)

    if show_percentages:
        _draw_percentage_labels(
            result,
            regions,
            selected_indices,
            (
                percentage_labels
                if percentage_labels is not None
                else _percentage_labels(image.size, regions)
            ),
            percentage_font_size,
            label_positions,
            text_colors,
        )

    if crop_to_regions:
        bounds = coverage.point(lambda value: 255 if value else 0).getbbox()
        if bounds is not None:
            margin = max((region.border_width + 1) // 2 for region in selected_regions)
            left, top, right, bottom = bounds
            bounds = (
                max(0, left - margin),
                max(0, top - margin),
                min(image.width, right + margin),
                min(image.height, bottom + margin),
            )
            result = result.crop(bounds)
    return result


def save_project_file(
    project_path: Path,
    image_path: Path,
    image: Image.Image,
    regions: list[Region],
    base_region_index: int | None,
    dpi: int,
    language: str,
) -> None:
    try:
        image_reference = os.path.relpath(image_path.resolve(), project_path.parent.resolve())
    except ValueError:
        image_reference = str(image_path.resolve())
    project = {
        "format": PROJECT_FORMAT,
        "version": PROJECT_VERSION,
        "image": {
            "path": image_reference,
            "width": image.width,
            "height": image.height,
        },
        "regions": [
            {
                "name": region.name,
                "points": [[x, y] for x, y in region.points],
                "color": region.color,
                "opacity": region.opacity,
                "border_width": region.border_width,
                "visible": region.visible,
                "percentage_base_index": region.percentage_base_index,
                "percentage_color": region.percentage_color,
                "percentage_label_position": (
                    list(region.percentage_label_position)
                    if region.percentage_label_position is not None
                    else None
                ),
            }
            for region in regions
        ],
        "base_region_index": base_region_index,
        "dpi": dpi,
        "language": language,
    }
    project_path.write_text(
        json.dumps(project, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def load_project_file(
    project_path: Path,
) -> tuple[Image.Image, Path, list[Region], int | None, int, str]:
    project = json.loads(project_path.read_text(encoding="utf-8"))
    if not isinstance(project, dict) or project.get("format") != PROJECT_FORMAT:
        raise ValueError("Not an Image Analyzer project file.")
    version = project.get("version")
    if not isinstance(version, int) or isinstance(version, bool) or version != PROJECT_VERSION:
        raise ValueError(f"Unsupported project version: {project.get('version')!r}")

    image_data = project.get("image")
    if not isinstance(image_data, dict) or not isinstance(image_data.get("path"), str):
        raise ValueError("Project does not contain a valid source image path.")
    image_width, image_height = image_data.get("width"), image_data.get("height")
    if (
        not isinstance(image_width, int)
        or isinstance(image_width, bool)
        or image_width <= 0
        or not isinstance(image_height, int)
        or isinstance(image_height, bool)
        or image_height <= 0
    ):
        raise ValueError("Project contains invalid source image dimensions.")
    image_path = Path(image_data["path"])
    if not image_path.is_absolute():
        image_path = project_path.parent / image_path
    image_path = image_path.resolve()
    with Image.open(image_path) as source:
        image = ImageOps.exif_transpose(source).convert("RGB")
    expected_size = (image_width, image_height)
    if expected_size != image.size:
        raise ValueError("The source image dimensions do not match the project.")

    region_data = project.get("regions")
    if not isinstance(region_data, list):
        raise ValueError("Project regions must be a list.")
    legacy_base_region_index = project.get("base_region_index")
    regions: list[Region] = []
    for item in region_data:
        if not isinstance(item, dict):
            raise ValueError("Project contains an invalid region.")
        name = item.get("name")
        color = item.get("color")
        points = item.get("points")
        opacity = item.get("opacity")
        border_width = item.get("border_width")
        visible = item.get("visible", True)
        percentage_base_index = item.get(
            "percentage_base_index", legacy_base_region_index
        )
        percentage_color = item.get("percentage_color")
        percentage_label_position = item.get("percentage_label_position")
        if (
            not isinstance(name, str)
            or not name.strip()
            or not isinstance(color, str)
            or re.fullmatch(r"#[0-9a-fA-F]{6}", color) is None
            or not isinstance(points, list)
            or not isinstance(opacity, int)
            or isinstance(opacity, bool)
            or not 0 <= opacity <= 100
            or not isinstance(border_width, int)
            or isinstance(border_width, bool)
            or not 1 <= border_width <= 12
            or not isinstance(visible, bool)
            or (
                percentage_base_index is not None
                and (
                    not isinstance(percentage_base_index, int)
                    or isinstance(percentage_base_index, bool)
                    or not 0 <= percentage_base_index < len(region_data)
                )
            )
            or (
                percentage_color is not None
                and (
                    not isinstance(percentage_color, str)
                    or re.fullmatch(r"#[0-9a-fA-F]{6}", percentage_color) is None
                )
            )
            or (
                percentage_label_position is not None
                and (
                    not isinstance(percentage_label_position, list)
                    or len(percentage_label_position) != 2
                    or any(
                        not isinstance(coordinate, (int, float))
                        or isinstance(coordinate, bool)
                        or not math.isfinite(coordinate)
                        or not 0 <= coordinate < dimension
                        for coordinate, dimension in zip(
                            percentage_label_position, (image.width, image.height)
                        )
                    )
                )
            )
        ):
            raise ValueError("Project contains invalid region settings.")
        parsed_points: list[tuple[float, float]] = []
        for point in points:
            if (
                not isinstance(point, list)
                or len(point) != 2
                or any(
                    not isinstance(coordinate, (int, float))
                    or isinstance(coordinate, bool)
                    or (
                        isinstance(coordinate, float)
                        and not math.isfinite(coordinate)
                    )
                    for coordinate in point
                )
            ):
                raise ValueError("Project contains invalid region coordinates.")
            x, y = point
            if not 0 <= x < image.width or not 0 <= y < image.height:
                raise ValueError("Project contains coordinates outside the source image.")
            parsed_points.append((float(x), float(y)))
        if len(parsed_points) < 3:
            raise ValueError("Project regions must contain at least three points.")
        regions.append(
            Region(
                name,
                parsed_points,
                color.lower(),
                opacity,
                border_width,
                visible,
                percentage_base_index,
                percentage_color.lower() if percentage_color else None,
                (
                    tuple(float(value) for value in percentage_label_position)
                    if percentage_label_position is not None
                    else None
                ),
            )
        )

    base_region_index = project.get("base_region_index")
    if base_region_index is not None and (
        not isinstance(base_region_index, int)
        or isinstance(base_region_index, bool)
        or not 0 <= base_region_index < len(regions)
    ):
        raise ValueError("Project contains an invalid base region.")
    dpi = project.get("dpi", 300)
    if not isinstance(dpi, int) or dpi not in {100, 200, 300, 400, 500, 600}:
        raise ValueError("Project contains an invalid DPI value.")
    language = project.get("language", "ru")
    if not isinstance(language, str) or language not in {"ru", "sr", "en"}:
        raise ValueError("Project contains an unsupported interface language.")
    return image, image_path, regions, base_region_index, dpi, language


class ImageAnalyzerApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.language = "ru"
        self._localized_widgets: list[tuple[tk.Widget, str]] = []
        self.root.title(self._t("title"))
        self.root.geometry("1280x820")
        self.root.minsize(850, 560)
        self.root.state("zoomed")

        self.image: Image.Image | None = None
        self.image_path: Path | None = None
        self.project_path: Path | None = None
        self.tk_image: ImageTk.PhotoImage | None = None
        self.regions: list[Region] = []
        self.selected_index: int | None = None
        self.selected_vertex: int | None = None
        self.dragging_vertex: int | None = None
        self.hovered_vertex: tuple[int, int] | None = None
        self.suppress_vertex_hover = False
        self.delete_vertex_mode = False
        self.drawing = False
        self.lasso_points: list[tuple[float, float]] = []
        self.lasso_region_index: int | None = None
        self.scale = 1.0
        self.zoom = 1.0
        self.pan_x = 0.0
        self.pan_y = 0.0
        self.pan_start: tuple[float, float, float, float] | None = None
        self.image_x = 0.0
        self.image_y = 0.0
        self.scale_x = 1.0
        self.scale_y = 1.0
        self.display_size = (0, 0)
        self.region_areas: list[int] = []
        self.union_area = 0
        self.base_region_index: int | None = None
        self.base_area_var = tk.StringVar(value=self._t("whole_image"))
        self.base_area_options: dict[str, int | None] = {}
        self.region_base_vars: list[tk.StringVar] = []
        self.region_base_options: list[dict[str, int | None]] = []
        self.region_row_labels: list[tk.Label] = []
        self.language_var = tk.StringVar(value="Русский")
        self.opacity_var = tk.IntVar(value=60)
        self.border_width_var = tk.IntVar(value=2)
        self.opacity_value_var = tk.StringVar(value="60%")
        self.border_width_value_var = tk.StringVar(value="2 px")
        self.dpi_var = tk.StringVar(value="300")
        self.status_var = tk.StringVar(value=self._t("status_load_image"))

        self._build_ui()
        self.canvas.bind("<Configure>", self._on_canvas_resize)
        self.canvas.bind("<ButtonPress-1>", self._on_press)
        self.canvas.bind("<B1-Motion>", self._on_drag)
        self.canvas.bind("<ButtonRelease-1>", self._on_release)
        self.canvas.bind("<Double-Button-1>", self._on_double_click)
        self.canvas.bind("<Motion>", self._on_mouse_move)
        self.canvas.bind("<MouseWheel>", self._on_mouse_wheel)
        self.canvas.bind("<Button-4>", self._on_mouse_wheel_linux)
        self.canvas.bind("<Button-5>", self._on_mouse_wheel_linux)
        self.canvas.bind("<ButtonPress-3>", self._on_pan_start)
        self.canvas.bind("<B3-Motion>", self._on_pan_drag)
        self.canvas.bind("<ButtonRelease-3>", self._on_pan_end)
        self.root.bind("<Delete>", self._delete_selected_region)
        self.root.bind("<Escape>", self._on_escape)

    def _build_ui(self) -> None:
        toolbar = ttk.Frame(self.root, padding=(8, 8, 8, 4))
        toolbar.pack(side=tk.TOP, fill=tk.X)
        self._button(toolbar, "load_image", command=self.load_image).pack(
            side=tk.LEFT, padx=(0, 6)
        )
        self._button(toolbar, "open_project", command=self.open_project).pack(
            side=tk.LEFT, padx=6
        )
        self._button(toolbar, "save_project", command=self.save_project).pack(
            side=tk.LEFT, padx=6
        )
        self._button(toolbar, "new_lasso", command=self.start_lasso).pack(
            side=tk.LEFT, padx=6
        )
        self.delete_region_button = self._button(
            toolbar, "delete_region", command=self.delete_selected_region
        )
        self.delete_region_button.pack(side=tk.LEFT, padx=6)
        self._button(toolbar, "export_image", command=self.export_image).pack(
            side=tk.LEFT, padx=6
        )
        self._button(toolbar, "export_excel", command=self.export_excel).pack(
            side=tk.LEFT, padx=6
        )
        self._label(toolbar, "dpi").pack(side=tk.LEFT, padx=(16, 4))
        ttk.Combobox(
            toolbar,
            textvariable=self.dpi_var,
            values=DPI_CHOICES,
            state="readonly",
            width=5,
        ).pack(side=tk.LEFT)
        self._label(toolbar, "language").pack(side=tk.RIGHT, padx=(8, 4))
        self.language_combo = ttk.Combobox(
            toolbar,
            textvariable=self.language_var,
            values=tuple(LANGUAGES),
            state="readonly",
            width=10,
        )
        self.language_combo.pack(side=tk.RIGHT)
        self.language_combo.bind("<<ComboboxSelected>>", self._on_language_change)

        body = ttk.Panedwindow(self.root, orient=tk.HORIZONTAL)
        body.pack(fill=tk.BOTH, expand=True, padx=8, pady=6)

        side = ttk.Frame(body, width=320, padding=8)
        body.add(side, weight=0)
        self._label(side, "regions", font=("", 11, "bold")).pack(
            anchor=tk.W, pady=(0, 5)
        )
        self._label(side, "region_bases_hint").pack(anchor=tk.W, pady=(0, 4))
        region_list_container = ttk.Frame(side)
        region_list_container.pack(fill=tk.BOTH, expand=True)
        self.region_list_canvas = tk.Canvas(
            region_list_container, highlightthickness=0, height=210
        )
        region_scrollbar = ttk.Scrollbar(
            region_list_container,
            orient=tk.VERTICAL,
            command=self.region_list_canvas.yview,
        )
        self.region_list_canvas.configure(yscrollcommand=region_scrollbar.set)
        region_scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.region_list_canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.region_list = ttk.Frame(self.region_list_canvas)
        self.region_list_window = self.region_list_canvas.create_window(
            (0, 0), window=self.region_list, anchor=tk.NW
        )
        self.region_list.bind(
            "<Configure>",
            lambda _event: self.region_list_canvas.configure(
                scrollregion=self.region_list_canvas.bbox("all")
            ),
        )
        self.region_list_canvas.bind(
            "<Configure>",
            lambda event: self.region_list_canvas.itemconfigure(
                self.region_list_window, width=event.width
            ),
        )
        self.region_list.bind("<Escape>", self._on_escape)
        self.region_list_canvas.bind("<Escape>", self._on_escape)

        region_buttons = ttk.Frame(side)
        region_buttons.pack(fill=tk.X, pady=(6, 10))
        self.rename_region_button = self._button(
            region_buttons, "rename", command=self.rename_region
        )
        self.rename_region_button.pack(side=tk.LEFT, padx=(0, 5))
        self.choose_color_button = self._button(
            region_buttons, "color", command=self.choose_color
        )
        self.choose_color_button.pack(side=tk.LEFT, padx=5)
        self.delete_vertex_button = self._button(
            region_buttons, "delete_vertex", command=self.toggle_delete_vertex_mode
        )
        self.delete_vertex_button.pack(side=tk.LEFT, padx=5)

        self._label(side, "fill_transparency").pack(anchor=tk.W)
        opacity_control_row = ttk.Frame(side)
        opacity_control_row.pack(fill=tk.X, pady=(3, 10))
        self.opacity_scale = ttk.Scale(
            opacity_control_row,
            from_=10,
            to=95,
            variable=self.opacity_var,
            command=self._on_opacity_change,
        )
        self.opacity_scale.pack(side=tk.LEFT, fill=tk.X, expand=True)
        ttk.Label(
            opacity_control_row,
            textvariable=self.opacity_value_var,
            width=5,
            anchor=tk.E,
        ).pack(side=tk.RIGHT, padx=(8, 0))
        self._label(side, "border_width").pack(anchor=tk.W)
        border_width_control_row = ttk.Frame(side)
        border_width_control_row.pack(fill=tk.X, pady=(3, 10))
        self.border_width_scale = ttk.Scale(
            border_width_control_row,
            from_=1,
            to=12,
            variable=self.border_width_var,
            command=self._on_border_width_change,
        )
        self.border_width_scale.pack(side=tk.LEFT, fill=tk.X, expand=True)
        ttk.Label(
            border_width_control_row,
            textvariable=self.border_width_value_var,
            width=6,
            anchor=tk.E,
        ).pack(side=tk.RIGHT, padx=(8, 0))
        self._sync_region_controls()

        ttk.Separator(side).pack(fill=tk.X, pady=6)
        self._label(side, "area_summary", font=("", 11, "bold")).pack(
            anchor=tk.W, pady=(0, 5)
        )
        self.summary_label = ttk.Label(
            side, text=self._t("summary_empty")
        )
        self.summary_label.pack(anchor=tk.W)
        self.stats_label = ttk.Label(side, text="", justify=tk.LEFT)
        self.stats_label.pack(anchor=tk.W, pady=(8, 0))

        canvas_frame = ttk.Frame(body)
        body.add(canvas_frame, weight=1)
        self.canvas = tk.Canvas(canvas_frame, background="#252525", highlightthickness=0)
        self.horizontal_scrollbar = ttk.Scrollbar(
            canvas_frame,
            orient=tk.HORIZONTAL,
            command=lambda *args: self._scroll_view("x", *args),
        )
        self.vertical_scrollbar = ttk.Scrollbar(
            canvas_frame,
            orient=tk.VERTICAL,
            command=lambda *args: self._scroll_view("y", *args),
        )
        self.canvas.grid(row=0, column=0, sticky="nsew")
        self.vertical_scrollbar.grid(row=0, column=1, sticky="ns")
        self.horizontal_scrollbar.grid(row=1, column=0, sticky="ew")
        ttk.Frame(canvas_frame).grid(row=1, column=1)
        canvas_frame.rowconfigure(0, weight=1)
        canvas_frame.columnconfigure(0, weight=1)

        ttk.Label(
            self.root,
            textvariable=self.status_var,
            anchor=tk.W,
            relief=tk.SUNKEN,
            padding=(8, 4),
        ).pack(side=tk.BOTTOM, fill=tk.X)

    def _t(self, key: str, **values: object) -> str:
        return translate(getattr(self, "language", "ru"), key, **values)

    def _label(self, parent: tk.Misc, key: str, **options: object) -> ttk.Label:
        widget = ttk.Label(parent, text=self._t(key), **options)
        self._localized_widgets.append((widget, key))
        return widget

    def _button(self, parent: tk.Misc, key: str, **options: object) -> ttk.Button:
        widget = ttk.Button(parent, text=self._t(key), **options)
        self._localized_widgets.append((widget, key))
        return widget

    def _on_language_change(self, _event: tk.Event | None = None) -> None:
        self.language = LANGUAGES[self.language_var.get()]
        self.apply_language()

    def apply_language(self) -> None:
        self.root.title(self._t("title"))
        for widget, key in self._localized_widgets:
            widget.configure(text=self._t(key))
        self._update_delete_vertex_button()
        self.status_var.set(self._t("language_changed"))
        self._refresh_regions()
        self._update_statistics()
        self._render()

    def _update_delete_vertex_button(self) -> None:
        if hasattr(self, "delete_vertex_button"):
            key = "cancel_delete" if self.delete_vertex_mode else "delete_vertex"
            self.delete_vertex_button.configure(text=self._t(key))

    def load_image(self) -> None:
        path = filedialog.askopenfilename(
            title=self._t("select_image"),
            filetypes=[
                (self._t("image_files"), "*.png *.jpg *.jpeg *.bmp *.tif *.tiff *.webp"),
                (self._t("all_files"), "*.*"),
            ],
        )
        if not path:
            return
        try:
            with Image.open(path) as source:
                loaded = ImageOps.exif_transpose(source).convert("RGB")
        except (OSError, ValueError) as exc:
            messagebox.showerror(self._t("open_image_error"), str(exc))
            return

        self.image = loaded
        self.image_path = Path(path)
        self.project_path = None
        self.regions.clear()
        self.delete_vertex_mode = False
        self._update_delete_vertex_button()
        self.selected_index = None
        self.selected_vertex = None
        self._sync_region_controls()
        self.base_region_index = None
        self.drawing = False
        self.lasso_points.clear()
        self.lasso_region_index = None
        self.hovered_vertex = None
        self.zoom = 1.0
        self.pan_x = 0.0
        self.pan_y = 0.0
        self.pan_start = None
        self.canvas.configure(cursor="arrow")
        self._refresh_regions()
        self._update_statistics()
        self._render()
        self.status_var.set(
            self._t(
                "image_loaded",
                name=self.image_path.name,
                width=loaded.width,
                height=loaded.height,
            )
        )

    def save_project(self) -> None:
        if self.image is None or self.image_path is None:
            messagebox.showinfo(self._t("no_image"), self._t("load_first"))
            return
        path = filedialog.asksaveasfilename(
            title=self._t("save_project_title"),
            defaultextension=".json",
            filetypes=[("JSON", "*.json")],
        )
        if not path:
            return
        project_path = Path(path)
        try:
            save_project_file(
                project_path,
                self.image_path,
                self.image,
                self.regions,
                None,
                int(self.dpi_var.get()),
                self.language,
            )
        except (OSError, ValueError, TypeError) as exc:
            messagebox.showerror(self._t("project_save_error"), str(exc))
            return
        self.project_path = project_path
        self.status_var.set(
            self._t("project_saved", name=project_path.name)
        )

    def open_project(self) -> None:
        path = filedialog.askopenfilename(
            title=self._t("open_project_title"),
            filetypes=[("JSON", "*.json"), (self._t("all_files"), "*.*")],
        )
        if not path:
            return
        project_path = Path(path)
        try:
            (
                image,
                image_path,
                regions,
                base_region_index,
                dpi,
                language,
            ) = load_project_file(project_path)
        except OSError as exc:
            messagebox.showerror(self._t("project_open_error"), str(exc))
            return
        except (ValueError, TypeError, KeyError):
            messagebox.showerror(
                self._t("project_open_error"), self._t("project_invalid")
            )
            return

        self.image = image
        self.image_path = image_path
        self.project_path = project_path
        self.regions = regions
        self.base_region_index = base_region_index
        self.dpi_var.set(str(dpi))
        self.language = language
        self.language_var.set(next(name for name, code in LANGUAGES.items() if code == language))
        self.delete_vertex_mode = False
        self._update_delete_vertex_button()
        self.selected_index = None
        self.selected_vertex = None
        self.dragging_vertex = None
        self.hovered_vertex = None
        self.drawing = False
        self.lasso_points.clear()
        self.lasso_region_index = None
        self.zoom = 1.0
        self.pan_x = 0.0
        self.pan_y = 0.0
        self.pan_start = None
        self.canvas.configure(cursor="arrow")
        self._sync_region_controls()
        self._update_statistics()
        self.apply_language()
        self.status_var.set(
            self._t(
                "project_opened",
                name=project_path.name,
                image=image_path.name,
            )
        )

    def start_lasso(self) -> None:
        if self.image is None:
            messagebox.showinfo(self._t("no_image"), self._t("load_first"))
            return
        self._set_delete_vertex_mode(False)
        if self.drawing:
            self._finish_lasso()
        self._sync_region_list_selection()
        self.selected_index = None
        self._sync_region_controls()
        self.suppress_vertex_hover = False
        self.drawing = True
        self.lasso_points.clear()
        self.lasso_region_index = None
        self.hovered_vertex = None
        self.selected_vertex = None
        self.dragging_vertex = None
        self.canvas.configure(cursor="crosshair")
        self.status_var.set(self._t("lasso_instructions"))
        self._render()

    def _on_canvas_resize(self, _event: tk.Event) -> None:
        if self.image is not None:
            self._render()

    def _fit_image(self) -> None:
        assert self.image is not None
        canvas_width = max(1, self.canvas.winfo_width())
        canvas_height = max(1, self.canvas.winfo_height())
        fit_scale = min(canvas_width / self.image.width, canvas_height / self.image.height)
        self.scale = fit_scale * self.zoom
        display_width = max(1, round(self.image.width * self.scale))
        display_height = max(1, round(self.image.height * self.scale))
        self.display_size = (display_width, display_height)
        self.scale_x = display_width / self.image.width
        self.scale_y = display_height / self.image.height
        self.image_x = (canvas_width - display_width) / 2 + self.pan_x
        self.image_y = (canvas_height - display_height) / 2 + self.pan_y

    def _render(self) -> None:
        self.canvas.delete("all")
        if self.image is None:
            if hasattr(self, "horizontal_scrollbar"):
                self.horizontal_scrollbar.set(0, 1)
                self.vertical_scrollbar.set(0, 1)
            self.canvas.create_text(
                max(1, self.canvas.winfo_width()) / 2,
                max(1, self.canvas.winfo_height()) / 2,
                text=self._t("canvas_empty"),
                fill="white",
                font=("", 16),
            )
            return

        self._fit_image()
        self._update_scrollbars()
        preview_source = self.image.resize(self.display_size, Image.Resampling.LANCZOS)
        scale_x = self.display_size[0] / self.image.width
        scale_y = self.display_size[1] / self.image.height
        preview_regions = [
            Region(
                region.name,
                [(x * scale_x, y * scale_y) for x, y in region.points],
                region.color,
                region.opacity,
                max(1, round(region.border_width * (scale_x + scale_y) / 2)),
                region.visible,
            )
            for region in self.regions
        ]
        preview = make_annotated_image(preview_source, preview_regions)
        self.tk_image = ImageTk.PhotoImage(preview)
        self.canvas.create_image(self.image_x, self.image_y, image=self.tk_image, anchor=tk.NW)

        if (
            self.drawing
            and self.lasso_points
            and (
                self.lasso_region_index is None
                or self.regions[self.lasso_region_index].visible
            )
        ):
            display_points = [self._image_to_canvas(point) for point in self.lasso_points]
            if len(display_points) > 1:
                if len(display_points) >= 3:
                    display_points.append(display_points[0])
                self.canvas.create_line(
                    *[coord for point in display_points for coord in point],
                    fill="#ffffff",
                    width=2,
                    dash=(5, 3),
                )
            for x, y in display_points:
                self.canvas.create_oval(
                    x - 4, y - 4, x + 4, y + 4, fill="#ffe600", outline="#101010", width=1
                )

        if (
            self.selected_index is not None
            and self.selected_index < len(self.regions)
            and self.regions[self.selected_index].visible
        ):
            region = self.regions[self.selected_index]
            for index, point in enumerate(region.points):
                x, y = self._image_to_canvas(point)
                is_hovered = self.hovered_vertex == (self.selected_index, index)
                is_selected = index == self.selected_vertex
                radius = 7 if is_hovered or is_selected else 5
                self.canvas.create_oval(
                    x - radius,
                    y - radius,
                    x + radius,
                    y + radius,
                    fill="#ffffff" if is_hovered or is_selected else region.color,
                    outline="#101010",
                    width=2,
                )

    def _update_scrollbars(self) -> None:
        if not hasattr(self, "horizontal_scrollbar"):
            return
        canvas_width = max(1, self.canvas.winfo_width())
        canvas_height = max(1, self.canvas.winfo_height())
        display_width, display_height = self.display_size

        if display_width <= canvas_width:
            self.horizontal_scrollbar.set(0, 1)
        else:
            first = max(0.0, min(1.0, -self.image_x / display_width))
            last = min(1.0, first + canvas_width / display_width)
            self.horizontal_scrollbar.set(first, last)

        if display_height <= canvas_height:
            self.vertical_scrollbar.set(0, 1)
        else:
            first = max(0.0, min(1.0, -self.image_y / display_height))
            last = min(1.0, first + canvas_height / display_height)
            self.vertical_scrollbar.set(first, last)

    def _scroll_view(self, axis: str, command: str, value: str, *extra: str) -> None:
        if self.image is None or axis not in {"x", "y"}:
            return
        horizontal = axis == "x"
        display_size = self.display_size[0 if horizontal else 1]
        viewport_size = max(
            1,
            self.canvas.winfo_width() if horizontal else self.canvas.winfo_height(),
        )
        if display_size <= viewport_size:
            return

        if command == "moveto":
            first = max(0.0, min(1.0, float(value)))
            first = min(first, 1.0 - viewport_size / display_size)
            origin = -first * display_size
            centered_origin = (viewport_size - display_size) / 2
            pan = origin - centered_origin
        elif command == "scroll":
            amount = float(value)
            step = viewport_size * (0.9 if extra and extra[0] == "pages" else 0.1)
            pan = (self.pan_x if horizontal else self.pan_y) - amount * step
        else:
            return

        if horizontal:
            self.pan_x = pan
        else:
            self.pan_y = pan
        self._render()

    def _image_to_canvas(self, point: tuple[float, float]) -> tuple[float, float]:
        return self.image_x + point[0] * self.scale_x, self.image_y + point[1] * self.scale_y

    def _canvas_to_image(self, x: float, y: float) -> tuple[float, float] | None:
        if self.image is None:
            return None
        image_x = (x - self.image_x) / self.scale_x
        image_y = (y - self.image_y) / self.scale_y
        if not (0 <= image_x < self.image.width and 0 <= image_y < self.image.height):
            return None
        return image_x, image_y

    def _on_mouse_wheel(self, event: tk.Event) -> None:
        if self.image is None or event.delta == 0:
            return
        self._zoom_at(event.x, event.y, 1.15 ** (event.delta / 120))

    def _on_mouse_wheel_linux(self, event: tk.Event) -> None:
        if event.num == 4:
            self._zoom_at(event.x, event.y, 1.15)
        elif event.num == 5:
            self._zoom_at(event.x, event.y, 1 / 1.15)

    def _zoom_at(self, x: float, y: float, factor: float) -> None:
        if self.image is None:
            return
        image_x = (x - self.image_x) / self.scale_x
        image_y = (y - self.image_y) / self.scale_y
        self.zoom = max(0.1, min(20.0, self.zoom * factor))
        self._fit_image()
        self.pan_x += x - (self.image_x + image_x * self.scale_x)
        self.pan_y += y - (self.image_y + image_y * self.scale_y)
        self._render()

    def _on_pan_start(self, event: tk.Event) -> None:
        if self.image is None:
            return
        self.pan_start = (event.x, event.y, self.pan_x, self.pan_y)
        self._set_canvas_cursor()

    def _on_pan_drag(self, event: tk.Event) -> None:
        if self.pan_start is None:
            return
        start_x, start_y, start_pan_x, start_pan_y = self.pan_start
        self.pan_x = start_pan_x + event.x - start_x
        self.pan_y = start_pan_y + event.y - start_y
        self._render()

    def _on_pan_end(self, _event: tk.Event) -> None:
        self.pan_start = None
        self._set_canvas_cursor()

    def _on_press(self, event: tk.Event) -> None:
        if self.delete_vertex_mode:
            nearest = (
                self._find_nearby_vertex(event.x, event.y, self.selected_index)
                if self.selected_index is not None
                and self.regions[self.selected_index].visible
                else None
            )
            if nearest is None:
                self.hovered_vertex = None
                self.selected_vertex = None
                self.status_var.set(
                    self._t("click_vertex_to_delete")
                )
                self._render()
                return
            self.hovered_vertex = nearest
            self.selected_index, self.selected_vertex = nearest
            self._sync_region_list_selection()
            point_count = len(self.regions[self.selected_index].points)
            self.delete_selected_vertex()
            if self.delete_vertex_mode and len(
                self.regions[self.selected_index].points
            ) < point_count:
                self.status_var.set(
                    self._t("vertex_deleted_next")
                )
            return
        if self.drawing:
            nearest = (
                self._find_nearby_vertex(event.x, event.y, self.lasso_region_index)
                if self.lasso_region_index is not None
                and self.regions[self.lasso_region_index].visible
                else None
            )
            if nearest is not None:
                self.selected_index, self.selected_vertex = nearest
                self.hovered_vertex = nearest
                self.dragging_vertex = nearest[1]
                self._sync_border_width_control()
                self._sync_region_controls()
                self._set_canvas_cursor()
                self._render()
                return
            point = self._canvas_to_image(event.x, event.y)
            if point is not None:
                if self.lasso_points:
                    last_x, last_y = self._image_to_canvas(self.lasso_points[-1])
                    if math.hypot(event.x - last_x, event.y - last_y) <= 3:
                        return
                self._add_lasso_point(point)
                self._render()
            return
        self.dragging_vertex = None
        nearest = (
            self._find_nearby_vertex(event.x, event.y, self.selected_index)
            if self.selected_index is not None
            and self.regions[self.selected_index].visible
            else None
        )
        self.hovered_vertex = nearest
        if nearest is not None:
            self.selected_vertex = nearest[1]
            self.dragging_vertex = self.selected_vertex
        elif self.selected_index is not None:
            point = self._canvas_to_image(event.x, event.y)
            if point is not None and self._insert_vertex_near_edge(point):
                return
            self.selected_vertex = None
        else:
            self.selected_vertex = None
        self._render()

    def _insert_vertex_near_edge(self, point: tuple[float, float]) -> bool:
        if self.selected_index is None:
            return False
        region = self.regions[self.selected_index]
        if len(region.points) < 3:
            return False

        closest_edge: tuple[float, int, tuple[float, float]] | None = None
        for index, start in enumerate(region.points):
            end = region.points[(index + 1) % len(region.points)]
            dx, dy = end[0] - start[0], end[1] - start[1]
            length_squared = dx * dx + dy * dy
            t = 0 if length_squared == 0 else max(
                0,
                min(
                    1,
                    ((point[0] - start[0]) * dx + (point[1] - start[1]) * dy)
                    / length_squared,
                ),
            )
            projected = (start[0] + t * dx, start[1] + t * dy)
            distance = math.hypot(
                (point[0] - projected[0]) * self.scale_x,
                (point[1] - projected[1]) * self.scale_y,
            )
            if closest_edge is None or distance < closest_edge[0]:
                closest_edge = (distance, index + 1, projected)

        if closest_edge is None or closest_edge[0] > VERTEX_MAGNET_RADIUS:
            return False
        _, insert_at, projected = closest_edge
        region.points.insert(insert_at, projected)
        self.selected_vertex = insert_at
        self.hovered_vertex = (self.selected_index, insert_at)
        self._update_statistics()
        self._refresh_regions()
        self._render()
        return True

    def _on_drag(self, event: tk.Event) -> None:
        point = self._canvas_to_image(event.x, event.y)
        if (
            self.dragging_vertex is not None
            and point is not None
            and self.selected_index is not None
        ):
            self.regions[self.selected_index].points[self.dragging_vertex] = point
            if self.drawing and self.selected_index == self.lasso_region_index:
                self.lasso_points = self.regions[self.selected_index].points
            self._render()
            return
        if self.drawing:
            if point is not None and self.lasso_points:
                last = self._image_to_canvas(self.lasso_points[-1])
                if math.hypot(event.x - last[0], event.y - last[1]) >= 3:
                    self._add_lasso_point(point)
                    self._render()
            return

    def _on_release(self, _event: tk.Event) -> None:
        if self.dragging_vertex is not None:
            self.dragging_vertex = None
            self._update_statistics()
            self._refresh_regions()
            hover_region_index = (
                self.lasso_region_index if self.drawing else self.selected_index
            )
            can_hover = self.drawing or (
                self.selected_index is not None
                and self.regions[self.selected_index].visible
            )
            self.hovered_vertex = (
                self._find_nearby_vertex(
                    _event.x, _event.y, hover_region_index
                )
                if can_hover
                else None
            )
            if self.hovered_vertex is not None:
                self.selected_vertex = self.hovered_vertex[1]
            else:
                self.selected_vertex = None
            self._set_canvas_cursor()
            self._render()
        elif self.drawing:
            return

    def _on_double_click(self, event: tk.Event) -> None:
        if self.drawing:
            if self.dragging_vertex is not None:
                self.dragging_vertex = None
                self._finish_lasso()
                return
            point = self._canvas_to_image(event.x, event.y)
            if point is not None and self.lasso_points:
                last = self._image_to_canvas(self.lasso_points[-1])
                if math.hypot(event.x - last[0], event.y - last[1]) > 3:
                    self._add_lasso_point(point)
            self._finish_lasso()
            return
        if self.selected_index is None:
            return
        region = self.regions[self.selected_index]
        if not region.visible:
            return
        point = self._canvas_to_image(event.x, event.y)
        if point is None or len(region.points) < 3:
            return
        if self._find_nearby_vertex(
            event.x, event.y, region_index=self.selected_index
        ) is not None:
            return
        if self._insert_vertex_near_edge(point):
            return

    def _on_mouse_move(self, event: tk.Event) -> None:
        if self.image is None or self.dragging_vertex is not None:
            return
        if getattr(self, "suppress_vertex_hover", False):
            nearby_after_escape = self._find_nearby_vertex(event.x, event.y)
            if nearby_after_escape is not None:
                return
            self.suppress_vertex_hover = False
        active_region_index = (
            self.lasso_region_index
            if self.drawing
            else (
                self.selected_index
                if self.selected_index is not None
                and self.regions[self.selected_index].visible
                else None
            )
        )
        nearby = (
            self._find_nearby_vertex(
                event.x, event.y, active_region_index
            )
            if active_region_index is not None
            else None
        )
        if nearby == self.hovered_vertex:
            return
        self.hovered_vertex = nearby
        if nearby is not None:
            self.selected_index, self.selected_vertex = nearby
            self._sync_region_list_selection()
            self._sync_opacity_control()
            self._sync_border_width_control()
            if self.delete_vertex_mode:
                self.status_var.set(
                    self._t("vertex_click_delete")
                )
            else:
                self.status_var.set(
                    self._t("vertex_drag")
                )
        else:
            self.selected_vertex = None
            if self.delete_vertex_mode:
                self.status_var.set(
                    self._t("delete_mode_hover")
                )
            else:
                self.status_var.set(self._t("hover_vertex_edit"))
            self._sync_region_controls()
        self._set_canvas_cursor()
        self._render()

    def _find_nearby_vertex(
        self, x: float, y: float, region_index: int | None = None
    ) -> tuple[int, int] | None:
        closest: tuple[float, int, int] | None = None
        for current_region_index, region in enumerate(self.regions):
            if region_index is not None and current_region_index != region_index:
                continue
            for vertex_index, point in enumerate(region.points):
                canvas_x, canvas_y = self._image_to_canvas(point)
                distance = math.hypot(x - canvas_x, y - canvas_y)
                if distance <= VERTEX_MAGNET_RADIUS and (
                    closest is None or distance < closest[0]
                ):
                    closest = (distance, current_region_index, vertex_index)
        if closest is None:
            return None
        return closest[1], closest[2]

    def _set_canvas_cursor(self) -> None:
        if self.pan_start is not None:
            cursor = "fleur"
        elif self.delete_vertex_mode:
            cursor = "crosshair"
        elif self.dragging_vertex is not None or self.hovered_vertex is not None:
            cursor = "fleur"
        elif self.drawing:
            cursor = "crosshair"
        else:
            cursor = "arrow"
        self.canvas.configure(cursor=cursor)

    def toggle_delete_vertex_mode(self) -> None:
        self._set_delete_vertex_mode(not self.delete_vertex_mode)

    def _set_delete_vertex_mode(self, enabled: bool) -> None:
        self.delete_vertex_mode = enabled
        if hasattr(self, "delete_vertex_button"):
            self.delete_vertex_button.configure(
                text=self._t("cancel_delete" if enabled else "delete_vertex")
            )
        if enabled:
            if self.drawing:
                self._finish_lasso()
            self.hovered_vertex = None
            self.selected_vertex = None
            self.status_var.set(
                self._t("delete_mode_hint")
            )
        else:
            self.hovered_vertex = None
            self.selected_vertex = None
            if hasattr(self, "status_var"):
                self.status_var.set(self._t("delete_mode_off"))
        if hasattr(self, "canvas"):
            self._set_canvas_cursor()
            self._render()

    def _add_lasso_point(self, point: tuple[float, float]) -> None:
        self.lasso_points.append(point)
        if len(self.lasso_points) == 3:
            number = len(self.regions) + 1
            existing_names = {region.name for region in self.regions}
            while self._t("region_name", number=number) in existing_names:
                number += 1
            region = Region(
                name=self._t("region_name", number=number),
                points=self.lasso_points,
                color=PALETTE[len(self.regions) % len(PALETTE)],
                opacity=LASSO_PREVIEW_OPACITY,
            )
            self.regions.append(region)
            self.lasso_region_index = len(self.regions) - 1
            self.selected_index = self.lasso_region_index
            self.selected_vertex = None
            self._sync_opacity_control()
            self._sync_border_width_control()
            self._sync_region_controls()
            self._sync_region_list_selection()
            self._update_statistics()
            self._refresh_regions()
        elif self.lasso_region_index is not None:
            self._update_statistics()
            self._refresh_regions()

    def _finish_lasso(self) -> None:
        self.drawing = False
        self.lasso_points = []
        if self.lasso_region_index is not None:
            region = self.regions[self.lasso_region_index]
            self.selected_index = self.lasso_region_index
            self.selected_vertex = None
            self._sync_opacity_control()
            self._sync_border_width_control()
            self._sync_region_controls()
            self._update_statistics()
            self._refresh_regions()
            self.status_var.set(self._t("added_region", name=region.name))
        else:
            self.status_var.set(self._t("region_not_created"))
        self.lasso_region_index = None
        self._set_canvas_cursor()
        self._render()

    def _on_region_select(self, _event: tk.Event | None = None) -> None:
        selected = self.region_list.curselection()
        self.selected_index = selected[0] if selected else None
        if (
            self.hovered_vertex is not None
            and self.hovered_vertex[0] != self.selected_index
        ):
            self.hovered_vertex = None
        if self.hovered_vertex is not None and self.hovered_vertex[0] == self.selected_index:
            self.selected_vertex = self.hovered_vertex[1]
        else:
            self.selected_vertex = None
        self.dragging_vertex = None
        if self.selected_index is not None:
            self._sync_opacity_control()
            self._sync_border_width_control()
        self._sync_region_controls()
        self._render()

    def _select_region_row(self, index: int) -> None:
        if not 0 <= index < len(self.regions):
            return
        if self.selected_index == index:
            if getattr(self, "delete_vertex_mode", False):
                self._set_delete_vertex_mode(False)
            self._clear_region_selection()
            return
        if getattr(self, "delete_vertex_mode", False):
            self._set_delete_vertex_mode(False)
        self.selected_index = index
        self.selected_vertex = None
        self.hovered_vertex = None
        self.dragging_vertex = None
        self._sync_region_list_selection()
        self._sync_opacity_control()
        self._sync_border_width_control()
        self._sync_region_controls()
        self._render()

    def _sync_region_list_selection(self) -> None:
        if hasattr(self.region_list, "curselection"):
            self.region_list.selection_clear(0, tk.END)
            if self.selected_index is not None:
                self.region_list.selection_set(self.selected_index)
                self.region_list.activate(self.selected_index)
        for index, label in enumerate(getattr(self, "region_row_labels", [])):
            label.configure(
                background="#dcecff"
                if index == self.selected_index
                else self.root.cget("background")
            )

    def _delete_selected_region_from_list(self, event: tk.Event) -> str:
        self.delete_selected_region()
        return "break"

    def _on_escape(self, _event: tk.Event | None = None) -> str:
        if self.delete_vertex_mode:
            self._set_delete_vertex_mode(False)
        self._clear_region_selection(suppress_hover=True)
        return "break"

    def _clear_region_selection(self, *, suppress_hover: bool = False) -> None:
        if self.drawing:
            self._finish_lasso()
        self.drawing = False
        self.lasso_points = []
        self.lasso_region_index = None
        self.selected_index = None
        self._sync_region_list_selection()
        self.selected_vertex = None
        self.hovered_vertex = None
        self.dragging_vertex = None
        self.suppress_vertex_hover = suppress_hover
        self._sync_region_controls()
        self._set_canvas_cursor()
        self._render()

    def _on_region_click(self, event: tk.Event) -> str:
        index = self.region_list.index(f"@{event.x},{event.y}")
        bounds = self.region_list.bbox(index) if 0 <= index < len(self.regions) else None
        if (
            bounds is None
            or event.y < bounds[1]
            or event.y >= bounds[1] + bounds[3]
            or not (0 <= index < len(self.regions))
        ):
            if self.delete_vertex_mode:
                self._set_delete_vertex_mode(False)
            self._clear_region_selection()
            return "break"

        if event.x <= 30:
            self._toggle_region_visibility(index)
            return "break"

        if index == self.selected_index:
            if self.delete_vertex_mode:
                self._set_delete_vertex_mode(False)
            self._clear_region_selection()
            return "break"

        self._select_region_row(index)
        return "break"

    def _toggle_region_visibility(self, index: int) -> None:
        if not 0 <= index < len(self.regions):
            return
        region = self.regions[index]
        region.visible = not region.visible
        self._refresh_regions()
        self._render()
        self.status_var.set(
            self._t(
                "region_visibility_shown" if region.visible else "region_visibility_hidden",
                name=region.name,
            )
        )

    def _refresh_regions(self) -> None:
        for child in self.region_list.winfo_children():
            child.destroy()
        self.region_row_labels = []
        self.region_base_vars = []
        self.region_base_options = []
        self.base_area_options = {self._t("whole_image"): None}
        for index, region in enumerate(self.regions):
            self.base_area_options[
                self._t(
                    "base_region_option", number=index + 1, name=region.name
                )
            ] = index
        for index, region in enumerate(self.regions):
            area = self.region_areas[index] if index < len(self.region_areas) else 0
            percentage_base = self._region_base_area_pixels(region)
            percent = calculate_area_percentage(area, percentage_base)
            percent_text = "—" if percent is None else f"{percent:.2f}%"

            row = ttk.Frame(self.region_list)
            row.pack(fill=tk.X, pady=1)
            eye_icon = "👁" if region.visible else "👁̸"
            ttk.Button(
                row,
                text=eye_icon,
                width=3,
                command=lambda region_index=index: self._toggle_region_visibility(
                    region_index
                ),
            ).pack(side=tk.LEFT, padx=(0, 3))
            row_label = tk.Label(
                row,
                text=f"{region.name} — {percent_text}",
                anchor=tk.W,
                foreground=region.color if region.visible else "#888888",
                cursor="hand2",
            )
            row_label.pack(side=tk.LEFT, fill=tk.X, expand=True)
            row_label.bind(
                "<Button-1>",
                lambda _event, region_index=index: self._select_region_row(
                    region_index
                ),
            )
            row_label.bind(
                "<Delete>",
                lambda event: self._delete_selected_region_from_list(event),
            )
            self.region_row_labels.append(row_label)

            choices = {self._t("whole_image"): None}
            for base_index, base_region in enumerate(self.regions):
                choices[
                    self._t(
                        "base_region_option",
                        number=base_index + 1,
                        name=base_region.name,
                    )
                ] = base_index
            self.region_base_options.append(choices)
            base_var = tk.StringVar(value=self._region_base_area_label(region))
            self.region_base_vars.append(base_var)
            base_row = ttk.Frame(row)
            base_row.pack(side=tk.RIGHT, padx=(4, 0))
            ttk.Label(base_row, text=self._t("region_base_area_prompt")).pack(
                side=tk.LEFT, padx=(0, 3)
            )
            base_combo = ttk.Combobox(
                base_row,
                textvariable=base_var,
                values=tuple(choices),
                state="readonly",
                width=15,
            )
            base_combo.pack(side=tk.LEFT)
            base_combo.bind(
                "<<ComboboxSelected>>",
                lambda _event, region_index=index: self._on_region_base_change(
                    region_index
                ),
            )
            row_label.bind("<Escape>", self._on_escape)

        self._sync_region_list_selection()

    def _update_statistics(self) -> None:
        if self.image is None:
            self.region_areas = []
            self.union_area = 0
            if hasattr(self, "summary_label"):
                self.summary_label.configure(text=self._t("summary_empty"))
            if hasattr(self, "stats_label"):
                self.stats_label.configure(text="")
            return
        self.region_areas, self.union_area = calculate_region_areas(
            self.image.size, self.regions
        )
        base_area = self._base_area_pixels()
        total_percent = calculate_area_percentage(sum(self.region_areas), base_area)
        union_percent = calculate_area_percentage(self.union_area, base_area)
        total_text = "—" if total_percent is None else f"{total_percent:.2f}%"
        union_text = "—" if union_percent is None else f"{union_percent:.2f}%"
        self.summary_label.configure(
            text="\n".join(
                (
                    self._t(
                        "summary_base", name=self._base_area_label(), area=base_area
                    ),
                    self._t("summary_sum", value=total_text),
                    self._t("summary_union", value=union_text),
                )
            )
        )
        self.stats_label.configure(
            text="\n".join(
                self._t(
                    "region_stat",
                    name=region.name,
                    percentage=self._format_percentage(
                        area, self._region_base_area_pixels(region)
                    ),
                    area=area,
                    base=self._region_base_area_label(region),
                )
                for region, area in zip(self.regions, self.region_areas)
            )
        )

    @staticmethod
    def _format_percentage(area: int, base_area: int) -> str:
        percentage = calculate_area_percentage(area, base_area)
        return "—" if percentage is None else f"{percentage:.2f}%"

    def _base_area_pixels(self) -> int:
        if self.image is None:
            return 0
        return self.image.width * self.image.height

    def _base_area_label(self) -> str:
        return self._t("whole_image")

    def _region_base_area_pixels(self, region: Region) -> int:
        if self.image is None:
            return 0
        index = region.percentage_base_index
        if index is None:
            return self.image.width * self.image.height
        if not 0 <= index < len(self.region_areas):
            return 0
        return self.region_areas[index]

    def _region_base_area_label(self, region: Region) -> str:
        choices = {self._t("whole_image"): None}
        for index, base_region in enumerate(self.regions):
            choices[
                self._t("base_region_option", number=index + 1, name=base_region.name)
            ] = index
        return next(
            (
                label
                for label, index in choices.items()
                if index == region.percentage_base_index
            ),
            self._t("whole_image"),
        )

    def _on_base_area_change(self, _event: tk.Event | None = None) -> None:
        self.base_region_index = self.base_area_options.get(self.base_area_var.get())
        self._update_statistics()
        self._refresh_regions()

    def _on_region_base_change(self, region_index: int) -> None:
        if not 0 <= region_index < len(self.regions):
            return
        self.regions[region_index].percentage_base_index = (
            self.region_base_options[region_index].get(
                self.region_base_vars[region_index].get()
            )
        )
        self._update_statistics()
        self._refresh_regions()

    def _on_opacity_change(self, _value: str) -> None:
        if self.selected_index is None or not (0 <= self.selected_index < len(self.regions)):
            return
        transparency = max(10, min(95, round(float(_value))))
        self.regions[self.selected_index].opacity = 100 - transparency
        if hasattr(self, "opacity_value_var"):
            self.opacity_value_var.set(f"{transparency}%")
        self._render()

    def _sync_opacity_control(self) -> None:
        if self.selected_index is None or not (0 <= self.selected_index < len(self.regions)):
            return
        self.opacity_var.set(100 - self.regions[self.selected_index].opacity)
        if hasattr(self, "opacity_value_var"):
            self.opacity_value_var.set(
                f"{100 - self.regions[self.selected_index].opacity}%"
            )

    def _on_border_width_change(self, value: str) -> None:
        if self.selected_index is None or not (0 <= self.selected_index < len(self.regions)):
            return
        self.regions[self.selected_index].border_width = max(
            1, min(12, round(float(value)))
        )
        if hasattr(self, "border_width_value_var"):
            self.border_width_value_var.set(
                f"{self.regions[self.selected_index].border_width} px"
            )
        self._render()

    def _sync_border_width_control(self) -> None:
        if self.selected_index is None or not (0 <= self.selected_index < len(self.regions)):
            return
        self.border_width_var.set(self.regions[self.selected_index].border_width)
        if hasattr(self, "border_width_value_var"):
            self.border_width_value_var.set(
                f"{self.regions[self.selected_index].border_width} px"
            )

    def _sync_region_controls(self) -> None:
        enabled = self.selected_index is not None and 0 <= self.selected_index < len(
            self.regions
        )
        state = tk.NORMAL if enabled else tk.DISABLED
        for widget_name in (
            "delete_region_button",
            "rename_region_button",
            "choose_color_button",
            "delete_vertex_button",
            "opacity_scale",
            "border_width_scale",
        ):
            widget = getattr(self, widget_name, None)
            if widget is not None:
                widget.configure(state=state)
        if hasattr(self, "opacity_value_var"):
            self.opacity_value_var.set(
                "—"
                if not enabled
                else f"{100 - self.regions[self.selected_index].opacity}%"
            )
        if hasattr(self, "border_width_value_var"):
            self.border_width_value_var.set(
                "—"
                if not enabled
                else f"{self.regions[self.selected_index].border_width} px"
            )

    def choose_color(self) -> None:
        region = self._selected_region()
        if region is None:
            return
        color = colorchooser.askcolor(color=region.color, title=self._t("region_color"))[1]
        if color:
            region.color = color
            self._refresh_regions()
            self._render()

    def rename_region(self) -> None:
        region = self._selected_region()
        if region is None:
            return
        name = simpledialog.askstring(
            self._t("rename_region"),
            self._t("name"),
            initialvalue=region.name,
            parent=self.root,
        )
        if name and name.strip():
            region.name = name.strip()
            self._refresh_regions()
            self._update_statistics()

    def delete_selected_region(self, _event: tk.Event | None = None) -> None:
        if self.selected_index is None:
            return
        if self.drawing:
            self._finish_lasso()
        removed_index = self.selected_index
        del self.regions[removed_index]
        for region in self.regions:
            if region.percentage_base_index == removed_index:
                region.percentage_base_index = None
            elif (
                region.percentage_base_index is not None
                and region.percentage_base_index > removed_index
            ):
                region.percentage_base_index -= 1
        if self.base_region_index == removed_index:
            self.base_region_index = None
        elif self.base_region_index is not None and self.base_region_index > removed_index:
            self.base_region_index -= 1
        self.selected_index = (
            min(removed_index, len(self.regions) - 1) if self.regions else None
        )
        self.selected_vertex = None
        self.dragging_vertex = None
        self.hovered_vertex = None
        if self.selected_index is None and self.delete_vertex_mode:
            self._set_delete_vertex_mode(False)
        self._sync_region_controls()
        self._update_statistics()
        self._refresh_regions()
        self._render()
        self.status_var.set(self._t("region_deleted"))

    def _delete_selected_region(self, event: tk.Event) -> None:
        if event.widget is self.region_list or event.widget in getattr(
            self, "region_row_labels", []
        ):
            self.delete_selected_region()

    def delete_selected_vertex(self) -> None:
        target = self.hovered_vertex
        if target is None and self.selected_index is not None and self.selected_vertex is not None:
            target = (self.selected_index, self.selected_vertex)
        if target is None:
            messagebox.showinfo(
                self._t("vertex_not_selected"), self._t("select_vertex")
            )
            return
        region_index, vertex_index = target
        if not (0 <= region_index < len(self.regions)):
            self.hovered_vertex = None
            self.selected_vertex = None
            return
        region = self.regions[region_index]
        if len(region.points) <= 3:
            messagebox.showinfo(
                self._t("minimum_points"), self._t("keep_three_points")
            )
            return
        del region.points[vertex_index]
        if self.drawing and region_index == self.lasso_region_index:
            self.lasso_points = region.points
        self.selected_index = region_index
        self.selected_vertex = None
        self.hovered_vertex = None
        self.dragging_vertex = None
        self._update_statistics()
        self._refresh_regions()
        self._set_canvas_cursor()
        self._render()

    def _selected_region(self) -> Region | None:
        if self.selected_index is None:
            messagebox.showinfo(
                self._t("region_not_selected"), self._t("select_region")
            )
            return None
        return self.regions[self.selected_index]

    def _make_export_preview(
        self,
        selected_indices: list[int],
        include_intersections: bool,
        crop_to_regions: bool,
        preview_size: tuple[int, int],
        show_percentages: bool = False,
        percentage_font_size: int = 24,
        label_positions: dict[int, tuple[float, float]] | None = None,
        text_colors: dict[int, str] | None = None,
    ) -> Image.Image:
        assert self.image is not None
        preview_width, preview_height = preview_size
        scale = min(
            preview_width / self.image.width,
            preview_height / self.image.height,
            1.0,
        )
        size = (
            max(1, round(self.image.width * scale)),
            max(1, round(self.image.height * scale)),
        )
        scale_x = size[0] / self.image.width
        scale_y = size[1] / self.image.height
        preview_source = self.image.resize(size, Image.Resampling.LANCZOS)
        preview_regions = [
            Region(
                region.name,
                [(x * scale_x, y * scale_y) for x, y in region.points],
                region.color,
                region.opacity,
                max(1, round(region.border_width * (scale_x + scale_y) / 2)),
                region.visible,
                region.percentage_base_index,
                region.percentage_color,
                (
                    (region.percentage_label_position[0] * scale_x,
                     region.percentage_label_position[1] * scale_y)
                    if region.percentage_label_position is not None
                    else None
                ),
            )
            for region in self.regions
        ]
        preview_positions = {
            index: (x * scale_x, y * scale_y)
            for index, (x, y) in (label_positions or {}).items()
        }
        return make_region_export(
            preview_source,
            preview_regions,
            selected_indices,
            include_intersections=include_intersections,
            crop_to_regions=crop_to_regions,
            show_percentages=show_percentages,
            percentage_font_size=max(1, round(percentage_font_size * (scale_x + scale_y) / 2)),
            percentage_labels=_percentage_labels(self.image.size, self.regions),
            label_positions=preview_positions,
            text_colors=text_colors,
        )

    def _choose_export_regions(
        self,
    ) -> tuple[
        list[int],
        bool,
        bool,
        bool,
        int,
        dict[int, tuple[float, float]],
        dict[int, str],
    ] | None:
        if not self.regions:
            return [], True, False, False, 24, {}, {}

        dialog = tk.Toplevel(self.root)
        dialog.title(self._t("export_options_title"))
        dialog.transient(self.root)
        dialog.geometry("1120x720")
        dialog.minsize(900, 560)
        dialog.grab_set()
        frame = ttk.Frame(dialog, padding=16)
        frame.pack(fill=tk.BOTH, expand=True)
        frame.columnconfigure(1, weight=1)
        frame.rowconfigure(0, weight=1)

        controls = ttk.Frame(frame)
        controls.grid(row=0, column=0, sticky="nsew", padx=(0, 16))
        ttk.Label(controls, text=self._t("export_regions_prompt")).pack(anchor=tk.W)

        selected_vars = [
            tk.BooleanVar(value=region.visible) for region in self.regions
        ]
        text_colors = {
            index: region.percentage_color or region.color
            for index, region in enumerate(self.regions)
        }
        label_positions = {
            index: region.percentage_label_position
            for index, region in enumerate(self.regions)
            if region.percentage_label_position is not None
        }
        regions_frame = ttk.Frame(controls)
        regions_frame.pack(fill=tk.X, pady=(8, 12))
        color_buttons: list[ttk.Button] = []

        def choose_text_color(region_index: int) -> None:
            selected_color = colorchooser.askcolor(
                color=text_colors[region_index],
                title=self._t("percentage_color"),
                parent=dialog,
            )[1]
            if selected_color is not None:
                text_colors[region_index] = selected_color.lower()
                color_buttons[region_index].configure(
                    text=f"{self._t('percentage_color')} {selected_color}"
                )
                schedule_preview()

        for index, (region, variable) in enumerate(zip(self.regions, selected_vars)):
            ttk.Checkbutton(
                regions_frame,
                text=region.name,
                variable=variable,
                state=tk.NORMAL if region.visible else tk.DISABLED,
            ).grid(row=index, column=0, sticky=tk.W, pady=2)
            color_button = ttk.Button(
                regions_frame,
                text=(
                    f"{self._t('percentage_color')} "
                    f"{region.percentage_color or region.color}"
                ),
                command=lambda region_index=index: choose_text_color(region_index),
                state=tk.NORMAL if region.visible else tk.DISABLED,
            )
            color_button.grid(row=index, column=1, sticky=tk.E, padx=(6, 0), pady=2)
            color_buttons.append(color_button)

        selection_buttons = ttk.Frame(controls)
        selection_buttons.pack(fill=tk.X, pady=(0, 10))

        include_intersections = tk.BooleanVar(value=True)
        crop_to_regions = tk.BooleanVar(value=True)
        show_percentages = tk.BooleanVar(value=False)
        font_size_var = tk.StringVar(value="24")
        preview_canvas = tk.Canvas(
            frame,
            background="#252525",
            highlightthickness=1,
            highlightbackground="#a0a0a0",
        )
        preview_canvas.grid(row=0, column=1, sticky="nsew")
        preview_canvas.create_text(
            350,
            250,
            text=self._t("export_preview"),
            fill="white",
            font=("", 14),
        )
        preview_photo: ImageTk.PhotoImage | None = None
        preview_after_id: str | None = None
        preview_state: dict[str, object] = {}

        def update_preview() -> None:
            nonlocal preview_photo
            width = max(1, preview_canvas.winfo_width() - 16)
            height = max(1, preview_canvas.winfo_height() - 16)
            selected_indices = [
                index
                for index, variable in enumerate(selected_vars)
                if variable.get()
            ]
            font_size = int(font_size_var.get())
            preview = self._make_export_preview(
                selected_indices,
                include_intersections.get(),
                crop_to_regions.get(),
                (width, height),
                show_percentages.get(),
                font_size,
                label_positions,
                text_colors,
            )
            scale = min(width / self.image.width, height / self.image.height, 1.0)
            source_size = (
                max(1, round(self.image.width * scale)),
                max(1, round(self.image.height * scale)),
            )
            scale_x = source_size[0] / self.image.width
            scale_y = source_size[1] / self.image.height
            preview_regions = [
                Region(
                    region.name,
                    [(x * scale_x, y * scale_y) for x, y in region.points],
                    region.color,
                    region.opacity,
                    max(1, round(region.border_width * (scale_x + scale_y) / 2)),
                    region.visible,
                )
                for region in self.regions
            ]
            crop_bounds = _region_export_crop_bounds(
                source_size, preview_regions, selected_indices, crop_to_regions.get()
            )
            preview_state.clear()
            preview_state.update(
                {
                    "scale_x": scale_x,
                    "scale_y": scale_y,
                    "crop_bounds": crop_bounds,
                    "preview_regions": preview_regions,
                    "font_size": max(
                        1, round(font_size * (scale_x + scale_y) / 2)
                    ),
                    "selected_indices": selected_indices,
                    "labels": _percentage_labels(self.image.size, self.regions),
                    "image_x": (preview_canvas.winfo_width() - preview.width) / 2,
                    "image_y": (preview_canvas.winfo_height() - preview.height) / 2,
                }
            )
            preview_photo = ImageTk.PhotoImage(preview)
            preview_canvas.delete("all")
            preview_canvas.create_image(
                preview_canvas.winfo_width() / 2,
                preview_canvas.winfo_height() / 2,
                image=preview_photo,
                anchor=tk.CENTER,
            )
            preview_canvas.image = preview_photo

        def on_label_press(event: tk.Event) -> None:
            if not show_percentages.get() or not preview_state:
                return
            bounds = preview_state.get("crop_bounds")
            if bounds is None:
                return
            crop_left, crop_top, _, _ = bounds
            local_x = event.x - float(preview_state["image_x"]) + crop_left
            local_y = event.y - float(preview_state["image_y"]) + crop_top
            scale_x = float(preview_state["scale_x"])
            scale_y = float(preview_state["scale_y"])
            regions = preview_state["preview_regions"]
            labels = preview_state["labels"]
            font = _percentage_font(int(preview_state["font_size"]))
            draw = ImageDraw.Draw(Image.new("RGB", (1, 1)))
            for index in reversed(preview_state["selected_indices"]):
                if not self.regions[index].visible or labels[index] is None:
                    continue
                region = regions[index]
                position = (
                    (label_positions.get(index) or self.regions[index].percentage_label_position)
                )
                if position is not None:
                    position = (position[0] * scale_x, position[1] * scale_y)
                else:
                    position = _default_percentage_label_position(
                        (
                            max(1, round(self.image.width * scale_x)),
                            max(1, round(self.image.height * scale_y)),
                        ),
                        region,
                        font,
                        labels[index],
                    )
                text_bounds = draw.textbbox(position, labels[index], font=font)
                preview_bounds = (
                    text_bounds[0] - crop_left,
                    text_bounds[1] - crop_top,
                    text_bounds[2] - crop_left,
                    text_bounds[3] - crop_top,
                )
                if (
                    preview_bounds[0] <= event.x - float(preview_state["image_x"]) <= preview_bounds[2]
                    and preview_bounds[1] <= event.y - float(preview_state["image_y"]) <= preview_bounds[3]
                ):
                    preview_state["dragging_label"] = index
                    return

        def on_label_drag(event: tk.Event) -> None:
            index = preview_state.get("dragging_label")
            bounds = preview_state.get("crop_bounds")
            if index is None or bounds is None:
                return
            crop_left, crop_top, _, _ = bounds
            label_positions[int(index)] = (
                max(
                    0.0,
                    min(
                        self.image.width - 1.0,
                        (event.x - float(preview_state["image_x"]) + crop_left)
                        / float(preview_state["scale_x"]),
                    ),
                ),
                max(
                    0.0,
                    min(
                        self.image.height - 1.0,
                        (event.y - float(preview_state["image_y"]) + crop_top)
                        / float(preview_state["scale_y"]),
                    ),
                ),
            )
            schedule_preview()

        def on_label_release(_event: tk.Event) -> None:
            preview_state.pop("dragging_label", None)

        def schedule_preview(*_args: object) -> None:
            nonlocal preview_after_id
            if preview_after_id is not None:
                dialog.after_cancel(preview_after_id)
            preview_after_id = dialog.after(80, update_preview)

        for variable in selected_vars:
            variable.trace_add("write", schedule_preview)

        def set_all(value: bool) -> None:
            for region, variable in zip(self.regions, selected_vars):
                if region.visible:
                    variable.set(value)

        ttk.Button(
            selection_buttons,
            text=self._t("select_all"),
            command=lambda: set_all(True),
        ).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(
            selection_buttons, text=self._t("select_none"), command=lambda: set_all(False)
        ).pack(side=tk.LEFT)

        ttk.Checkbutton(
            controls,
            text=self._t("include_intersections"),
            variable=include_intersections,
            command=schedule_preview,
        ).pack(anchor=tk.W)
        ttk.Checkbutton(
            controls,
            text=self._t("crop_export"),
            variable=crop_to_regions,
            command=schedule_preview,
        ).pack(anchor=tk.W, pady=(4, 12))
        ttk.Checkbutton(
            controls,
            text=self._t("show_percentages"),
            variable=show_percentages,
            command=schedule_preview,
        ).pack(anchor=tk.W, pady=(0, 6))
        ttk.Label(controls, text=self._t("percentage_font_size")).pack(anchor=tk.W)
        font_size_combo = ttk.Combobox(
            controls,
            textvariable=font_size_var,
            values=tuple(str(value) for value in range(12, 49, 4)),
            state="readonly",
            width=8,
        )
        font_size_combo.pack(anchor=tk.W, pady=(2, 10))
        font_size_var.trace_add("write", schedule_preview)
        preview_canvas.bind("<ButtonPress-1>", on_label_press)
        preview_canvas.bind("<B1-Motion>", on_label_drag)
        preview_canvas.bind("<ButtonRelease-1>", on_label_release)
        include_intersections.trace_add("write", schedule_preview)
        crop_to_regions.trace_add("write", schedule_preview)
        preview_canvas.bind("<Configure>", schedule_preview)

        result: dict[
            str,
            tuple[
                list[int],
                bool,
                bool,
                bool,
                int,
                dict[int, tuple[float, float]],
                dict[int, str],
            ]
            | None,
        ] = {"value": None}

        def accept() -> None:
            indices = [
                index for index, variable in enumerate(selected_vars) if variable.get()
            ]
            if not indices:
                messagebox.showinfo(
                    self._t("export_options_title"),
                    self._t("select_export_region"),
                    parent=dialog,
                )
                return
            result["value"] = (
                indices,
                include_intersections.get(),
                crop_to_regions.get(),
                show_percentages.get(),
                int(font_size_var.get()),
                label_positions,
                text_colors,
            )
            dialog.destroy()

        buttons = ttk.Frame(frame)
        buttons.grid(row=1, column=0, columnspan=2, sticky="ew", pady=(12, 0))
        ttk.Button(
            buttons, text=self._t("cancel"), command=dialog.destroy
        ).pack(side=tk.RIGHT)
        ttk.Button(
            buttons, text=self._t("continue"), command=accept
        ).pack(side=tk.RIGHT, padx=(0, 6))
        dialog.protocol("WM_DELETE_WINDOW", dialog.destroy)
        dialog.update_idletasks()
        update_preview()
        dialog.wait_window()
        return result["value"]

    def export_image(self) -> None:
        if self.image is None:
            messagebox.showinfo(self._t("no_image"), self._t("load_first"))
            return
        options = self._choose_export_regions()
        if options is None:
            return
        (
            selected_indices,
            include_intersections,
            crop_to_regions,
            show_percentages,
            percentage_font_size,
            label_positions,
            text_colors,
        ) = options
        path = filedialog.asksaveasfilename(
            title=self._t("image_export_title"),
            defaultextension=".png",
            filetypes=[("PNG", "*.png"), ("JPEG", "*.jpg *.jpeg"), ("TIFF", "*.tif *.tiff")],
        )
        if not path:
            return
        dpi = int(self.dpi_var.get())
        result = make_region_export(
            self.image,
            self.regions,
            selected_indices,
            include_intersections=include_intersections,
            crop_to_regions=crop_to_regions,
            show_percentages=show_percentages,
            percentage_font_size=percentage_font_size,
            percentage_labels=_percentage_labels(self.image.size, self.regions),
            label_positions=label_positions,
            text_colors=text_colors,
        )
        try:
            result.save(path, dpi=(dpi, dpi))
        except (OSError, ValueError) as exc:
            messagebox.showerror(self._t("export_error"), str(exc))
            return
        for index in selected_indices:
            self.regions[index].percentage_color = text_colors.get(
                index,
                self.regions[index].percentage_color or self.regions[index].color,
            )
            if index in label_positions:
                self.regions[index].percentage_label_position = label_positions[index]
        self.status_var.set(
            self._t(
                "image_saved",
                name=Path(path).name,
                width=result.width,
                height=result.height,
                dpi=dpi,
            )
        )

    def export_excel(self) -> None:
        if self.image is None:
            messagebox.showinfo(self._t("no_image"), self._t("load_first"))
            return
        path = filedialog.asksaveasfilename(
            title=self._t("table_export_title"),
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
        )
        if not path:
            return

        base_pixels = self._base_area_pixels()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = self._t("excel_sheet")
        sheet.append(
            [self._t("excel_image"), self.image_path.name if self.image_path else ""]
        )
        sheet.append(
            [self._t("excel_size"), f"{self.image.width} × {self.image.height}"]
        )
        sheet.append([self._t("excel_base_area"), base_pixels])
        sheet.append([self._t("excel_base"), self._base_area_label()])
        sheet.append([])
        sheet.append(
            [
                self._t("excel_region"),
                self._t("excel_region_area"),
                self._t("excel_region_base"),
                self._t("excel_base_share"),
                self._t("excel_color"),
            ]
        )
        for cell in sheet[6]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        for region, area in zip(self.regions, self.region_areas):
            region_base_pixels = self._region_base_area_pixels(region)
            percentage = calculate_area_percentage(area, region_base_pixels)
            sheet.append(
                [
                    region.name,
                    area,
                    self._region_base_area_label(region),
                    None if percentage is None else round(percentage, 4),
                    region.color,
                ]
            )
            sheet.cell(sheet.max_row, 5).fill = PatternFill(
                "solid", fgColor=region.color.lstrip("#")
            )

        sheet.append([])
        sheet.append([self._t("excel_sum_overlap_area"), sum(self.region_areas)])
        total_percentage = calculate_area_percentage(sum(self.region_areas), base_pixels)
        sheet.append(
            [
                self._t("excel_sum_overlap_percent"),
                None if total_percentage is None else round(total_percentage, 4),
            ]
        )
        sheet.append([self._t("excel_union_area"), self.union_area])
        union_percentage = calculate_area_percentage(self.union_area, base_pixels)
        sheet.append(
            [
                self._t("excel_union_percent"),
                None if union_percentage is None else round(union_percentage, 4),
            ]
        )
        sheet.column_dimensions["A"].width = 54
        sheet.column_dimensions["B"].width = 24
        sheet.column_dimensions["C"].width = 24
        sheet.column_dimensions["D"].width = 24
        sheet.column_dimensions["E"].width = 14
        try:
            workbook.save(path)
        except (OSError, ValueError) as exc:
            messagebox.showerror(self._t("export_error"), str(exc))
            return
        self.status_var.set(self._t("table_saved", name=Path(path).name))


def main() -> None:
    root = tk.Tk()
    ImageAnalyzerApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
