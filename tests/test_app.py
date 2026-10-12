import tempfile
import unittest
from pathlib import Path
from string import Formatter
from types import SimpleNamespace
from unittest.mock import Mock, patch

from PIL import Image
from openpyxl import load_workbook

from app import (
    LASSO_PREVIEW_OPACITY,
    Region,
    calculate_area_percentage,
    calculate_region_areas,
    make_annotated_image,
    make_region_export,
    save_project_file,
    load_project_file,
)
from translations import LANGUAGES, translate
import translations


class LocalizationTests(unittest.TestCase):
    def test_russian_is_default_and_three_language_choices_are_available(self) -> None:
        self.assertEqual(tuple(LANGUAGES), ("Русский", "Srpski", "English"))
        self.assertEqual(translate("ru", "title"), "Анализ областей на изображении")

    def test_user_interface_text_is_available_in_all_languages(self) -> None:
        self.assertEqual(translate("ru", "delete_vertex"), "Удалить точку")
        self.assertEqual(translate("sr", "delete_vertex"), "Obriši tačku")
        self.assertEqual(translate("en", "delete_vertex"), "Delete point")

    def test_all_translation_templates_format_in_every_language(self) -> None:
        formatter = Formatter()
        for key, localized_templates in translations._TEXTS.items():
            for language, template in zip(("ru", "sr", "en"), localized_templates):
                fields = {
                    field
                    for _, field, _, _ in formatter.parse(template)
                    if field is not None
                }
                values = {
                    field: 1234 if field in {"area", "width", "height", "dpi", "number"} else "value"
                    for field in fields
                }
                translate(language, key, **values)

    def test_region_labels_keep_numbers_and_custom_names_when_translated(self) -> None:
        self.assertEqual(
            translate("en", "base_region_option", number=2, name="Custom"),
            "Region 2: Custom",
        )
        self.assertEqual(
            translate("sr", "added_region", name="Custom"),
            "Dodata je oblast „Custom”.",
        )

    def test_switching_language_updates_window_controls_and_status(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.language = "sr"
        app.root = Mock()
        app.status_var = Mock()
        app.delete_vertex_mode = False
        app.delete_vertex_button = Mock()
        label = Mock()
        app._localized_widgets = [(label, "load_image")]
        app._refresh_regions = Mock()
        app._update_statistics = Mock()
        app._render = Mock()

        app.apply_language()

        app.root.title.assert_called_once_with("Analiza oblasti na slici")
        label.configure.assert_called_once_with(text="Učitaj sliku")
        app.status_var.set.assert_called_once_with(
            "Jezik interfejsa je promenjen."
        )
        app._refresh_regions.assert_called_once()
        app._update_statistics.assert_called_once()
        app._render.assert_called_once()


class RegionAreaTests(unittest.TestCase):
    def test_calculates_individual_and_union_areas(self) -> None:
        regions = [
            Region("A", [(0, 0), (4, 0), (4, 4), (0, 4)], "#ff0000"),
            Region("B", [(2, 2), (6, 2), (6, 6), (2, 6)], "#00ff00"),
        ]

        areas, union_area = calculate_region_areas((10, 10), regions)

        self.assertEqual(areas, [25, 25])
        self.assertEqual(union_area, 41)
        self.assertEqual(sum(areas), 50)

    def test_empty_region_list_has_zero_area(self) -> None:
        self.assertEqual(calculate_region_areas((10, 10), []), ([], 0))

    def test_area_percentage_uses_selected_base_and_handles_zero(self) -> None:
        self.assertEqual(calculate_area_percentage(25, 100), 25.0)
        self.assertEqual(calculate_area_percentage(150, 100), 150.0)
        self.assertIsNone(calculate_area_percentage(25, 0))

    def test_annotated_export_has_source_dimensions(self) -> None:
        source = Image.new("RGB", (12, 9), "white")
        region = Region("A", [(1, 1), (8, 1), (8, 6), (1, 6)], "#ff0000", 50)

        result = make_annotated_image(source, [region])

        self.assertEqual(result.size, source.size)
        self.assertNotEqual(result.getpixel((4, 3)), source.getpixel((4, 3)))

    def test_region_export_crops_to_selected_regions_and_preserves_source_background(self) -> None:
        source = Image.new("RGB", (24, 24), "#dddddd")
        region = Region(
            "Selected",
            [(8, 8), (16, 8), (16, 16), (8, 16)],
            "#00ff00",
            opacity=100,
            border_width=2,
        )

        result = make_region_export(source, [region], [0])

        self.assertEqual(result.size, (11, 11))
        self.assertEqual(result.getpixel((5, 5)), (0, 255, 0))

    def test_region_export_can_hide_overlap_fills_without_hiding_individual_regions(self) -> None:
        source = Image.new("RGB", (24, 24), "#dddddd")
        regions = [
            Region(
                "A",
                [(2, 2), (14, 2), (14, 14), (2, 14)],
                "#ff0000",
                opacity=100,
            ),
            Region(
                "B",
                [(9, 9), (20, 9), (20, 20), (9, 20)],
                "#00ff00",
                opacity=100,
            ),
        ]

        result = make_region_export(
            source, regions, [0, 1], include_intersections=False, crop_to_regions=False
        )

        self.assertEqual(result.getpixel((11, 11)), (221, 221, 221))
        self.assertEqual(result.getpixel((5, 5)), (255, 0, 0))
        self.assertEqual(result.getpixel((18, 18)), (0, 255, 0))


class ProjectFileTests(unittest.TestCase):
    def test_project_round_trip_preserves_regions_and_resolves_relative_image_path(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            image_path = root / "base.png"
            project_path = root / "projects" / "analysis.json"
            project_path.parent.mkdir()
            image = Image.new("RGB", (32, 24), "#eeeeee")
            image.save(image_path)
            regions = [
                Region(
                    "Measured area",
                    [(2.5, 3.0), (20.0, 3.0), (18.0, 17.5)],
                    "#aabbcc",
                    opacity=63,
                    border_width=5,
                )
            ]

            save_project_file(
                project_path, image_path, image, regions, 0, 500, "sr"
            )
            loaded = load_project_file(project_path)

            loaded_image, loaded_path, loaded_regions, base_index, dpi, language = loaded
            self.assertEqual(loaded_image.size, image.size)
            self.assertEqual(loaded_path, image_path.resolve())
            self.assertEqual(loaded_regions, regions)
            self.assertEqual(base_index, 0)
            self.assertEqual(dpi, 500)
            self.assertEqual(language, "sr")

    def test_invalid_project_version_is_rejected(self) -> None:
        import json

        with tempfile.TemporaryDirectory() as temp_dir:
            project_path = Path(temp_dir) / "invalid.json"
            project_path.write_text(
                json.dumps({"format": "image-analyzer-project", "version": 2}),
                encoding="utf-8",
            )

            with self.assertRaisesRegex(ValueError, "Unsupported project version"):
                load_project_file(project_path)


class AppBehaviorTests(unittest.TestCase):
    def test_starting_new_lasso_deselects_previous_region_until_new_region_exists(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.image = Image.new("RGB", (20, 20), "white")
        app.regions = [
            Region("Previous", [(1, 1), (10, 1), (10, 10)], "#ff0000")
        ]
        app.selected_index = 0
        app.selected_vertex = 1
        app.hovered_vertex = (0, 1)
        app.suppress_vertex_hover = False
        app.dragging_vertex = None
        app.delete_vertex_mode = False
        app.drawing = False
        app.pan_start = None
        app.lasso_points = []
        app.lasso_region_index = None
        app.region_list = Mock()
        app.canvas = Mock()
        app.status_var = Mock()
        app._sync_region_controls = Mock()
        app._render = Mock()

        app.start_lasso()

        self.assertIsNone(app.selected_index)
        self.assertIsNone(app.selected_vertex)
        self.assertIsNone(app.hovered_vertex)
        self.assertIsNone(app.dragging_vertex)
        self.assertTrue(app.drawing)
        app.region_list.selection_clear.assert_called_with(0, "end")
        app._sync_region_controls.assert_called()

    def test_outline_width_changes_rendered_border(self) -> None:
        source = Image.new("RGB", (30, 30), "white")
        points = [(5, 5), (24, 5), (24, 24), (5, 24)]

        thin = make_annotated_image(
            source, [Region("Thin", points, "#ff0000", border_width=1)]
        )
        thick = make_annotated_image(
            source, [Region("Thick", points, "#ff0000", border_width=6)]
        )

        self.assertGreater(
            thin.getpixel((7, 15))[1],
            thick.getpixel((7, 15))[1],
        )

    def test_border_width_slider_updates_selected_region(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.selected_index = 0
        app.regions = [Region("A", [(1, 1), (8, 1), (8, 8)], "#ff0000")]
        app.border_width_value_var = Mock()
        app._render = Mock()

        app._on_border_width_change("7")

        self.assertEqual(app.regions[0].border_width, 7)
        app.border_width_value_var.set.assert_called_once_with("7 px")
        app._render.assert_called_once()

    def test_setting_value_labels_match_selected_region_and_empty_state(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.regions = [
            Region("A", [(1, 1), (8, 1), (8, 8)], "#ff0000", opacity=25, border_width=5)
        ]
        app.selected_index = 0
        app.opacity_value_var = Mock()
        app.border_width_value_var = Mock()

        app._sync_region_controls()

        app.opacity_value_var.set.assert_called_once_with("75%")
        app.border_width_value_var.set.assert_called_once_with("5 px")

        app.opacity_value_var.set.reset_mock()
        app.border_width_value_var.set.reset_mock()
        app.selected_index = None
        app._sync_region_controls()

        app.opacity_value_var.set.assert_called_once_with("—")
        app.border_width_value_var.set.assert_called_once_with("—")

    def test_lasso_preview_uses_high_transparency(self) -> None:
        source = Image.new("RGB", (30, 30), "white")
        region = Region(
            "Preview",
            [(2, 2), (25, 2), (25, 25)],
            "#ff0000",
            opacity=LASSO_PREVIEW_OPACITY,
        )

        result = make_annotated_image(source, [region])

        red, green, blue = result.getpixel((10, 5))
        self.assertEqual(red, 255)
        self.assertGreater(green, 210)
        self.assertGreater(blue, 210)

    def test_transparency_levels_change_fill_alpha_in_expected_direction(self) -> None:
        source = Image.new("RGB", (30, 30), "white")
        points = [(2, 2), (25, 2), (25, 25)]

        nearly_opaque = make_annotated_image(
            source, [Region("A", points, "#ff0000", opacity=90)]
        )
        mostly_transparent = make_annotated_image(
            source, [Region("A", points, "#ff0000", opacity=10)]
        )

        opaque_pixel = nearly_opaque.getpixel((10, 5))
        transparent_pixel = mostly_transparent.getpixel((10, 5))
        self.assertLess(opaque_pixel[1], transparent_pixel[1])
        self.assertLess(opaque_pixel[2], transparent_pixel[2])

    def test_transparency_slider_updates_selected_region_and_redraws(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.selected_index = 0
        app.regions = [Region("A", [(1, 1), (8, 1), (8, 8)], "#ff0000", opacity=40)]
        app.opacity_value_var = Mock()
        app._render = Mock()

        app._on_opacity_change("85")

        self.assertEqual(app.regions[0].opacity, 15)
        app.opacity_value_var.set.assert_called_with("85%")
        app._render.assert_called_once()

        app._on_opacity_change("20")

        self.assertEqual(app.regions[0].opacity, 80)
        app.opacity_value_var.set.assert_called_with("20%")
        self.assertEqual(app._render.call_count, 2)

    def test_opacity_control_displays_transparency_for_selected_region(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.selected_index = 0
        app.regions = [Region("A", [(1, 1), (8, 1), (8, 8)], "#ff0000", opacity=15)]
        app.opacity_var = Mock()

        app._sync_opacity_control()

        app.opacity_var.set.assert_called_once_with(85)

    def test_image_export_preserves_pixels_and_records_dpi(self) -> None:
        from app import ImageAnalyzerApp

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "annotated.png"
            app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
            app.image = Image.new("RGB", (20, 12), "white")
            app.regions = []
            app.dpi_var = Mock()
            app.dpi_var.get.return_value = "200"
            app.status_var = Mock()

            with patch("app.filedialog.asksaveasfilename", return_value=str(output_path)):
                app.export_image()

            with Image.open(output_path) as exported:
                self.assertEqual(exported.size, (20, 12))
                self.assertAlmostEqual(exported.info["dpi"][0], 200, delta=1)

    def test_image_export_uses_selected_regions_and_crop_options(self) -> None:
        from app import ImageAnalyzerApp

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "selected.png"
            app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
            app.image = Image.new("RGB", (24, 24), "white")
            app.regions = [
                Region(
                    "A",
                    [(2, 2), (10, 2), (10, 10), (2, 10)],
                    "#ff0000",
                    opacity=100,
                ),
                Region(
                    "B",
                    [(8, 8), (16, 8), (16, 16), (8, 16)],
                    "#00ff00",
                    opacity=100,
                ),
            ]
            app._choose_export_regions = Mock(return_value=([1], False, True))
            app.dpi_var = Mock()
            app.dpi_var.get.return_value = "300"
            app.status_var = Mock()

            with patch("app.filedialog.asksaveasfilename", return_value=str(output_path)):
                app.export_image()

            with Image.open(output_path) as exported:
                self.assertEqual(exported.size, (11, 11))
                self.assertEqual(exported.getpixel((5, 5)), (0, 255, 0))

    def test_excel_export_includes_selected_base_and_area_percentages(self) -> None:
        from app import ImageAnalyzerApp

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "areas.xlsx"
            app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
            app.image = Image.new("RGB", (10, 10), "white")
            app.image_path = Path("source.png")
            app.regions = [
                Region("A", [(1, 1), (5, 1), (5, 5)], "#ff0000", opacity=40)
            ]
            app.region_areas = [25]
            app.union_area = 25
            app.base_region_index = None
            app.status_var = Mock()

            with patch("app.filedialog.asksaveasfilename", return_value=str(output_path)):
                app.export_excel()

            workbook = load_workbook(output_path, data_only=True)
            sheet = workbook["Области"]
            self.assertEqual(sheet["B3"].value, 100)
            self.assertEqual(sheet["B4"].value, "Всё изображение")
            self.assertEqual(sheet["A7"].value, "A")
            self.assertEqual(sheet["C7"].value, 25)
            workbook.close()

    def test_excel_export_uses_the_selected_language(self) -> None:
        from app import ImageAnalyzerApp

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "areas-en.xlsx"
            app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
            app.language = "en"
            app.image = Image.new("RGB", (10, 10), "white")
            app.image_path = Path("source.png")
            app.regions = []
            app.region_areas = []
            app.union_area = 0
            app.base_region_index = None
            app.status_var = Mock()

            with patch("app.filedialog.asksaveasfilename", return_value=str(output_path)):
                app.export_excel()

            workbook = load_workbook(output_path, data_only=True)
            sheet = workbook["Regions"]
            self.assertEqual(sheet["A1"].value, "Image")
            self.assertEqual(sheet["A6"].value, "Region")
            self.assertEqual(sheet["A8"].value, "Sum of regions (with overlaps), px²")
            workbook.close()

    def test_third_lasso_point_adds_region_and_later_points_update_it(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.regions = []
        app.lasso_points = []
        app.lasso_region_index = None
        app.selected_index = None
        app.selected_vertex = None
        app.opacity_var = Mock()
        app.region_list = Mock()
        app._update_statistics = Mock()
        app._refresh_regions = Mock()
        app._sync_region_controls = Mock()
        app._sync_border_width_control = Mock()

        points = [(1.0, 1.0), (8.0, 1.0), (8.0, 8.0), (1.0, 8.0)]
        for point in points[:3]:
            app._add_lasso_point(point)

        self.assertEqual(len(app.regions), 1)
        self.assertEqual(app.regions[0].name, "Область 1")
        self.assertIs(app.regions[0].points, app.lasso_points)
        app._sync_region_controls.assert_called_once()

        app._add_lasso_point(points[3])

        self.assertEqual(len(app.regions), 1)
        self.assertEqual(app.regions[0].points, points)
        self.assertEqual(app._refresh_regions.call_count, 2)

    def test_finishing_lasso_preserves_selected_transparency(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        region = Region("Active", [(1, 1), (8, 1), (8, 8)], "#ff0000", opacity=25)
        app.regions = [region]
        app.drawing = True
        app.lasso_points = region.points
        app.lasso_region_index = 0
        app.selected_index = 0
        app.selected_vertex = None
        app.border_width_var = Mock()
        app.opacity_var = Mock()
        app.status_var = Mock()
        app._update_statistics = Mock()
        app._refresh_regions = Mock()
        app._sync_border_width_control = Mock()
        app._set_canvas_cursor = Mock()
        app._render = Mock()

        app._finish_lasso()

        self.assertFalse(app.drawing)
        self.assertEqual(region.opacity, 25)
        app.opacity_var.set.assert_called_once_with(75)
        app._sync_border_width_control.assert_called_once()
        app._render.assert_called_once()

    def test_nearby_vertex_uses_magnet_radius_and_prefers_nearest(self) -> None:
        from app import ImageAnalyzerApp, VERTEX_MAGNET_RADIUS

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.regions = [
            Region("A", [(10.0, 10.0), (30.0, 30.0)], "#ff0000"),
            Region("B", [(11.0, 10.0)], "#00ff00"),
        ]
        app.image_x = 0.0
        app.image_y = 0.0
        app.scale = 1.0
        app.scale_x = 1.0
        app.scale_y = 1.0

        self.assertEqual(
            app._find_nearby_vertex(10.8, 10.0),
            (1, 0),
        )
        self.assertEqual(
            app._find_nearby_vertex(10.8, 10.0, region_index=0),
            (0, 0),
        )
        self.assertIsNone(
            app._find_nearby_vertex(30.0 + VERTEX_MAGNET_RADIUS + 1, 30.0)
        )

    def test_hover_magnet_activates_vertex_editing_in_active_lasso(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.image = Image.new("RGB", (40, 40), "white")
        app.image_x = 0.0
        app.image_y = 0.0
        app.scale = 1.0
        app.scale_x = 1.0
        app.scale_y = 1.0
        app.regions = [
            Region("Active", [(10.0, 10.0), (25.0, 10.0), (25.0, 25.0)], "#ff0000")
        ]
        app.drawing = True
        app.delete_vertex_mode = False
        app.dragging_vertex = None
        app.pan_start = None
        app.pan_start = None
        app.lasso_region_index = 0
        app.hovered_vertex = None
        app.selected_index = 0
        app.selected_vertex = None
        app.region_list = Mock()
        app.opacity_var = Mock()
        app._sync_border_width_control = Mock()
        app._sync_region_controls = Mock()
        app.status_var = Mock()
        app.canvas = Mock()
        app._render = Mock()

        event = SimpleNamespace(x=10.0, y=10.0)
        app._on_mouse_move(event)

        self.assertEqual(app.hovered_vertex, (0, 0))
        self.assertEqual(app.selected_vertex, 0)
        app.canvas.configure.assert_called_with(cursor="fleur")

        app._on_press(event)

        self.assertEqual(app.dragging_vertex, 0)

    def test_hovering_or_clicking_unselected_region_does_not_activate_it(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.image = Image.new("RGB", (80, 80), "white")
        app.image_x = 0.0
        app.image_y = 0.0
        app.scale = 1.0
        app.scale_x = 1.0
        app.scale_y = 1.0
        app.regions = [
            Region(
                "Selected",
                [(5.0, 5.0), (20.0, 5.0), (20.0, 20.0)],
                "#ff0000",
            ),
            Region(
                "Unselected",
                [(50.0, 50.0), (70.0, 50.0), (70.0, 70.0)],
                "#00ff00",
            ),
        ]
        app.drawing = False
        app.lasso_region_index = None
        app.delete_vertex_mode = False
        app.dragging_vertex = None
        app.pan_start = None
        app.hovered_vertex = None
        app.selected_index = 0
        app.selected_vertex = None
        app.region_list = Mock()
        app.canvas = Mock()
        app._sync_region_controls = Mock()
        app._sync_opacity_control = Mock()
        app._sync_border_width_control = Mock()
        app._set_canvas_cursor = Mock()
        app._render = Mock()

        event = SimpleNamespace(x=50.0, y=50.0)
        app._on_mouse_move(event)
        app._on_press(event)

        self.assertEqual(app.selected_index, 0)
        self.assertIsNone(app.hovered_vertex)
        self.assertIsNone(app.selected_vertex)
        app.region_list.selection_set.assert_not_called()
        app.region_list.activate.assert_not_called()

    def test_hover_during_new_lasso_does_not_select_existing_regions(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.image = Image.new("RGB", (40, 40), "white")
        app.image_x = 0.0
        app.image_y = 0.0
        app.scale = 1.0
        app.scale_x = 1.0
        app.scale_y = 1.0
        app.regions = [
            Region("First", [(10.0, 10.0), (25.0, 10.0), (25.0, 25.0)], "#ff0000"),
            Region("Second", [(5.0, 5.0), (30.0, 5.0), (30.0, 30.0)], "#00ff00"),
        ]
        app.drawing = True
        app.lasso_region_index = None
        app.dragging_vertex = None
        app.hovered_vertex = None
        app.selected_index = None
        app.selected_vertex = None
        app.delete_vertex_mode = False
        app.pan_start = None
        app.region_list = Mock()
        app.opacity_var = Mock()
        app._sync_border_width_control = Mock()
        app._sync_region_controls = Mock()
        app.status_var = Mock()
        app.canvas = Mock()
        app._render = Mock()

        app._on_mouse_move(SimpleNamespace(x=10.0, y=10.0))

        self.assertIsNone(app.selected_index)
        self.assertIsNone(app.hovered_vertex)
        self.assertIsNone(app.selected_vertex)
        app.region_list.selection_set.assert_not_called()
        app._render.assert_not_called()

    def test_single_click_near_selected_region_edge_inserts_editable_vertex(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.image = Image.new("RGB", (80, 80), "white")
        app.image_x = 0.0
        app.image_y = 0.0
        app.scale = 1.0
        app.scale_x = 1.0
        app.scale_y = 1.0
        app.regions = [
            Region(
                "Editable",
                [(10.0, 10.0), (50.0, 10.0), (50.0, 50.0), (10.0, 50.0)],
                "#ff0000",
            )
        ]
        app.selected_index = 0
        app.selected_vertex = None
        app.hovered_vertex = None
        app.dragging_vertex = None
        app.delete_vertex_mode = False
        app.drawing = False
        app.region_list = Mock()
        app._update_statistics = Mock()
        app._refresh_regions = Mock()
        app._render = Mock()
        app._sync_opacity_control = Mock()
        app._sync_border_width_control = Mock()
        app._sync_region_controls = Mock()

        app._on_press(SimpleNamespace(x=30.0, y=12.0))

        self.assertEqual(len(app.regions[0].points), 5)
        self.assertEqual(app.regions[0].points[1], (30.0, 10.0))
        self.assertEqual(app.selected_vertex, 1)
        app._update_statistics.assert_called_once()
        app._refresh_regions.assert_called_once()
        app._render.assert_called_once()

    def test_click_away_from_selected_region_edge_does_not_insert_vertex(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.image = Image.new("RGB", (80, 80), "white")
        app.image_x = 0.0
        app.image_y = 0.0
        app.scale = 1.0
        app.scale_x = 1.0
        app.scale_y = 1.0
        app.regions = [
            Region(
                "Editable",
                [(10.0, 10.0), (50.0, 10.0), (50.0, 50.0), (10.0, 50.0)],
                "#ff0000",
            )
        ]
        app.selected_index = 0
        app.selected_vertex = None
        app.hovered_vertex = None
        app.dragging_vertex = None
        app.delete_vertex_mode = False
        app.drawing = False
        app.region_list = Mock()
        app._update_statistics = Mock()
        app._refresh_regions = Mock()
        app._render = Mock()
        app._sync_opacity_control = Mock()
        app._sync_border_width_control = Mock()
        app._sync_region_controls = Mock()

        app._on_press(SimpleNamespace(x=30.0, y=30.0))

        self.assertEqual(len(app.regions[0].points), 4)
        app._update_statistics.assert_not_called()
        app._refresh_regions.assert_not_called()

    def test_double_click_on_new_vertex_does_not_insert_duplicate_point(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.image = Image.new("RGB", (80, 80), "white")
        app.image_x = 0.0
        app.image_y = 0.0
        app.scale = 1.0
        app.scale_x = 1.0
        app.scale_y = 1.0
        app.regions = [
            Region(
                "Editable",
                [(10.0, 10.0), (30.0, 10.0), (50.0, 10.0), (50.0, 50.0)],
                "#ff0000",
            )
        ]
        app.selected_index = 0
        app.drawing = False

        app._on_double_click(SimpleNamespace(x=30.0, y=10.0))

        self.assertEqual(len(app.regions[0].points), 4)

    def test_zoom_keeps_image_point_under_cursor_fixed(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.image = Image.new("RGB", (100, 100), "white")
        app.canvas = Mock()
        app.canvas.winfo_width.return_value = 200
        app.canvas.winfo_height.return_value = 100
        app.zoom = 1.0
        app.pan_x = 0.0
        app.pan_y = 0.0
        app.scale = 1.0
        app.image_x = 50.0
        app.image_y = 0.0
        app.scale_x = 1.0
        app.scale_y = 1.0
        app._render = app._fit_image

        cursor = (80.0, 40.0)
        image_point_before = (
            (cursor[0] - app.image_x) / app.scale_x,
            (cursor[1] - app.image_y) / app.scale_y,
        )

        app._zoom_at(*cursor, 1.15)

        image_point_after = (
            (cursor[0] - app.image_x) / app.scale_x,
            (cursor[1] - app.image_y) / app.scale_y,
        )
        self.assertAlmostEqual(app.zoom, 1.15)
        self.assertAlmostEqual(image_point_after[0], image_point_before[0])
        self.assertAlmostEqual(image_point_after[1], image_point_before[1])

    def test_mouse_wheel_zoom_uses_pointer_position_and_direction(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.image = Image.new("RGB", (20, 20), "white")
        app._zoom_at = Mock()

        app._on_mouse_wheel(SimpleNamespace(delta=120, x=45, y=60))

        app._zoom_at.assert_called_once_with(45, 60, 1.15)

        app._zoom_at.reset_mock()
        app._on_mouse_wheel(SimpleNamespace(delta=-120, x=30, y=40))

        app._zoom_at.assert_called_once_with(30, 40, 1 / 1.15)

    def test_scrollbar_moveto_updates_pan_and_renders(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.image = Image.new("RGB", (300, 240), "white")
        app.canvas = Mock()
        app.canvas.winfo_width.return_value = 100
        app.canvas.winfo_height.return_value = 80
        app.display_size = (300, 240)
        app.pan_x = 0.0
        app.pan_y = 0.0
        app._render = Mock()

        app._scroll_view("x", "moveto", "0.5")
        self.assertEqual(app.pan_x, -50.0)
        self.assertEqual(app.pan_y, 0.0)
        app._render.assert_called_once()

        app._scroll_view("y", "moveto", "0.5")
        self.assertEqual(app.pan_y, -40.0)
        self.assertEqual(app._render.call_count, 2)

    def test_scrollbar_thumb_tracks_visible_image_area(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.canvas = Mock()
        app.canvas.winfo_width.return_value = 100
        app.canvas.winfo_height.return_value = 80
        app.display_size = (300, 240)
        app.image_x = -100.0
        app.image_y = -80.0
        app.horizontal_scrollbar = Mock()
        app.vertical_scrollbar = Mock()

        app._update_scrollbars()

        app.horizontal_scrollbar.set.assert_called_once_with(
            100 / 300, 200 / 300
        )
        app.vertical_scrollbar.set.assert_called_once_with(1 / 3, 2 / 3)

    def test_right_drag_updates_pan_without_resetting_zoom(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.image = Image.new("RGB", (20, 20), "white")
        app.pan_x = 4.0
        app.pan_y = -3.0
        app.pan_start = (20.0, 30.0, app.pan_x, app.pan_y)
        app.canvas = Mock()
        app._render = Mock()
        app._set_canvas_cursor = Mock()

        app._on_pan_drag(SimpleNamespace(x=35.0, y=24.0))

        self.assertEqual((app.pan_x, app.pan_y), (19.0, -9.0))
        app._render.assert_called_once()

    def test_delete_selected_vertex_removes_vertex_by_hover_target(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.regions = [
            Region(
                "Editable",
                [(1.0, 1.0), (8.0, 1.0), (8.0, 8.0), (1.0, 8.0)],
                "#ff0000",
            )
        ]
        app.hovered_vertex = (0, 2)
        app.selected_index = 0
        app.selected_vertex = None
        app.dragging_vertex = None
        app.drawing = False
        app.canvas = Mock()
        app._update_statistics = Mock()
        app._refresh_regions = Mock()
        app._set_canvas_cursor = Mock()
        app._render = Mock()

        app.delete_selected_vertex()

        self.assertEqual(app.regions[0].points, [(1.0, 1.0), (8.0, 1.0), (1.0, 8.0)])
        self.assertIsNone(app.hovered_vertex)
        app._update_statistics.assert_called_once()
        app._render.assert_called_once()

    def test_delete_mode_removes_magnetically_targeted_vertex_on_left_click(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.image = Image.new("RGB", (30, 30), "white")
        app.image_x = 0.0
        app.image_y = 0.0
        app.scale = 1.0
        app.scale_x = 1.0
        app.scale_y = 1.0
        app.regions = [
            Region(
                "Editable",
                [(5.0, 5.0), (25.0, 5.0), (25.0, 25.0), (5.0, 25.0)],
                "#ff0000",
            )
        ]
        app.delete_vertex_mode = True
        app.hovered_vertex = None
        app.selected_index = 0
        app.selected_vertex = None
        app.dragging_vertex = None
        app.drawing = False
        app.lasso_region_index = None
        app.region_list = Mock()
        app.status_var = Mock()
        app.pan_start = None
        app.canvas = Mock()
        app._update_statistics = Mock()
        app._refresh_regions = Mock()
        app._render = Mock()

        app._on_press(SimpleNamespace(x=5.0, y=5.0))

        self.assertEqual(
            app.regions[0].points,
            [(25.0, 5.0), (25.0, 25.0), (5.0, 25.0)],
        )
        self.assertIsNone(app.hovered_vertex)
        app._update_statistics.assert_called_once()
        app._render.assert_called_once()
        app.status_var.set.assert_called_with(
            "Точка удалена. Наведите курсор на следующую вершину или отмените режим."
        )

    def test_delete_mode_requires_a_vertex_inside_magnet_radius(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.image = Image.new("RGB", (30, 30), "white")
        app.image_x = 0.0
        app.image_y = 0.0
        app.scale = 1.0
        app.scale_x = 1.0
        app.scale_y = 1.0
        app.regions = [
            Region(
                "Editable",
                [(5.0, 5.0), (25.0, 5.0), (25.0, 25.0), (5.0, 25.0)],
                "#ff0000",
            )
        ]
        app.delete_vertex_mode = True
        app.hovered_vertex = None
        app.selected_index = 0
        app.selected_vertex = None
        app.region_list = Mock()
        app.status_var = Mock()
        app.pan_start = None
        app._render = Mock()
        app.delete_selected_vertex = Mock()

        app._on_press(SimpleNamespace(x=15.0, y=15.0))

        app.delete_selected_vertex.assert_not_called()
        app.status_var.set.assert_called_with(
            "Подведите курсор к вершине и щёлкните левой кнопкой для удаления."
        )

    def test_statistics_use_selected_region_as_percentage_base(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.image = Image.new("RGB", (10, 10), "white")
        app.regions = [
            Region("Base", [], "#ff0000"),
            Region("Other", [], "#00ff00"),
        ]
        app.region_areas = [25, 10]
        app.union_area = 30
        app.base_region_index = 0
        app.summary_label = Mock()
        app.stats_label = Mock()

        with patch("app.calculate_region_areas", return_value=([25, 10], 30)):
            app._update_statistics()

        summary = app.summary_label.configure.call_args.kwargs["text"]
        stats = app.stats_label.configure.call_args.kwargs["text"]
        self.assertIn("Base (25 px²)", summary)
        self.assertIn("Сумма областей (с наложениями): 140.00%", summary)
        self.assertIn("Объединение (без наложений): 120.00%", summary)
        self.assertIn("Base: 100.00% (25 px²)", stats)
        self.assertIn("Other: 40.00% (10 px²)", stats)

    def test_clicking_active_region_in_list_deselects_it(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.regions = [
            Region("A", [(1, 1), (8, 1), (8, 8)], "#ff0000"),
            Region("B", [(2, 2), (9, 2), (9, 9)], "#00ff00"),
        ]
        app.selected_index = 0
        app.selected_vertex = 1
        app.hovered_vertex = (0, 1)
        app.dragging_vertex = None
        app.delete_vertex_mode = False
        app.pan_start = None
        app.region_list = Mock()
        app.region_list.index.return_value = 0
        app.region_list.bbox.return_value = (0, 0, 100, 20)
        app.canvas = Mock()
        app.drawing = False
        app._sync_region_controls = Mock()
        app._render = Mock()

        result = app._on_region_click(SimpleNamespace(x=10, y=10))

        self.assertEqual(result, "break")
        self.assertIsNone(app.selected_index)
        self.assertIsNone(app.selected_vertex)
        self.assertIsNone(app.hovered_vertex)
        app.region_list.selection_clear.assert_called_once_with(0, "end")
        app._sync_region_controls.assert_called_once()
        app._render.assert_called_once()

    def test_clicking_other_region_selects_it_instead_of_deselecting(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.regions = [
            Region("A", [(1, 1), (8, 1), (8, 8)], "#ff0000"),
            Region("B", [(2, 2), (9, 2), (9, 9)], "#00ff00"),
        ]
        app.selected_index = 0
        app.selected_vertex = None
        app.hovered_vertex = None
        app.dragging_vertex = None
        app.region_list = Mock()
        app.region_list.index.return_value = 1
        app.region_list.bbox.return_value = (0, 20, 100, 20)
        app.region_list.curselection.return_value = (1,)
        app._sync_opacity_control = Mock()
        app._sync_border_width_control = Mock()
        app._sync_region_controls = Mock()
        app._render = Mock()

        result = app._on_region_click(SimpleNamespace(x=10, y=30))

        self.assertEqual(result, "break")
        self.assertEqual(app.selected_index, 1)
        app.region_list.selection_set.assert_called_once_with(1)
        app._sync_opacity_control.assert_called_once()
        app._sync_border_width_control.assert_called_once()
        app._sync_region_controls.assert_called_once()
        app._render.assert_called_once()

    def test_region_controls_disable_without_selection_and_enable_with_one(self) -> None:
        import tkinter as tk
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.regions = [Region("A", [(1, 1), (8, 1), (8, 8)], "#ff0000")]
        controls = [Mock() for _ in range(6)]
        (
            app.delete_region_button,
            app.rename_region_button,
            app.choose_color_button,
            app.delete_vertex_button,
            app.opacity_scale,
            app.border_width_scale,
        ) = controls

        app.selected_index = None
        app._sync_region_controls()
        for control in controls:
            control.configure.assert_called_with(state=tk.DISABLED)

        for control in controls:
            control.configure.reset_mock()
        app.selected_index = 0
        app._sync_region_controls()
        for control in controls:
            control.configure.assert_called_with(state=tk.NORMAL)

    def test_escape_clears_area_and_vertex_selection(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.delete_vertex_mode = False
        app.selected_index = 0
        app.selected_vertex = 2
        app.hovered_vertex = (0, 2)
        app.dragging_vertex = 2
        app.region_list = Mock()
        app.canvas = Mock()
        app.pan_start = None
        app.drawing = False
        app._sync_region_controls = Mock()
        app._set_canvas_cursor = Mock()
        app._render = Mock()

        result = app._on_escape()

        self.assertEqual(result, "break")
        self.assertIsNone(app.selected_index)
        self.assertIsNone(app.selected_vertex)
        self.assertIsNone(app.hovered_vertex)
        self.assertIsNone(app.dragging_vertex)
        app.region_list.selection_clear.assert_called_once_with(0, "end")
        app._sync_region_controls.assert_called_once()
        app._render.assert_called_once()

    def test_escape_does_not_reselect_region_when_pointer_moves_over_vertex(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.image = Image.new("RGB", (40, 40), "white")
        app.image_x = 0.0
        app.image_y = 0.0
        app.scale = 1.0
        app.scale_x = 1.0
        app.scale_y = 1.0
        app.regions = [
            Region("First", [(10.0, 10.0), (25.0, 10.0), (25.0, 25.0)], "#ff0000"),
            Region("Second", [(5.0, 5.0), (30.0, 5.0), (30.0, 30.0)], "#00ff00"),
        ]
        app.delete_vertex_mode = False
        app.selected_index = 0
        app.selected_vertex = 0
        app.hovered_vertex = (0, 0)
        app.dragging_vertex = None
        app.drawing = False
        app.pan_start = None
        app.region_list = Mock()
        app.canvas = Mock()
        app.status_var = Mock()
        app._sync_opacity_control = Mock()
        app._sync_border_width_control = Mock()
        app._sync_region_controls = Mock()
        app._set_canvas_cursor = Mock()
        app._render = Mock()

        app._on_escape()
        app._on_mouse_move(SimpleNamespace(x=10.0, y=10.0))

        self.assertIsNone(app.selected_index)
        self.assertIsNone(app.hovered_vertex)
        self.assertTrue(app.suppress_vertex_hover)

        app._on_mouse_move(SimpleNamespace(x=39.0, y=39.0))
        self.assertFalse(app.suppress_vertex_hover)
        app._on_mouse_move(SimpleNamespace(x=10.0, y=10.0))
        self.assertIsNone(app.selected_index)
        self.assertIsNone(app.hovered_vertex)

    def test_clicking_second_region_uses_clicked_list_row_not_stale_vertex_hover(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.regions = [
            Region("First", [(1, 1), (8, 1), (8, 8)], "#ff0000"),
            Region("Second", [(2, 2), (9, 2), (9, 9)], "#00ff00"),
            Region("Third", [(3, 3), (10, 3), (10, 10)], "#0000ff"),
        ]
        app.selected_index = 0
        app.selected_vertex = 0
        app.hovered_vertex = (0, 0)
        app.dragging_vertex = None
        app.delete_vertex_mode = False
        app.pan_start = None
        app.region_list = Mock()
        app.region_list.index.return_value = 1
        app.region_list.bbox.return_value = (0, 20, 100, 20)
        app.region_list.curselection.return_value = (1,)
        app._sync_opacity_control = Mock()
        app._sync_border_width_control = Mock()
        app._sync_region_controls = Mock()
        app._render = Mock()

        app._on_region_click(SimpleNamespace(x=10, y=30))

        self.assertEqual(app.selected_index, 1)
        self.assertIsNone(app.hovered_vertex)
        self.assertIsNone(app.selected_vertex)
        app.region_list.index.assert_called_once_with("@10,30")

    def test_clicking_empty_space_below_regions_clears_selection(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.regions = [
            Region("First", [(1, 1), (8, 1), (8, 8)], "#ff0000"),
            Region("Second", [(2, 2), (9, 2), (9, 9)], "#00ff00"),
        ]
        app.selected_index = 1
        app.selected_vertex = 2
        app.hovered_vertex = (1, 2)
        app.dragging_vertex = None
        app.delete_vertex_mode = False
        app.drawing = False
        app.pan_start = None
        app.region_list = Mock()
        app.region_list.index.return_value = 1
        app.region_list.bbox.return_value = (0, 20, 100, 20)
        app._sync_region_controls = Mock()
        app._set_canvas_cursor = Mock()
        app._render = Mock()

        result = app._on_region_click(SimpleNamespace(x=10, y=75))

        self.assertEqual(result, "break")
        self.assertIsNone(app.selected_index)
        self.assertIsNone(app.selected_vertex)
        self.assertIsNone(app.hovered_vertex)
        app.region_list.selection_clear.assert_called_once_with(0, "end")
        app.region_list.selection_set.assert_not_called()
        app._sync_region_controls.assert_called_once()
        app._render.assert_called_once()

    def test_escape_finishes_active_lasso_and_hides_its_points(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        app.delete_vertex_mode = False
        app.drawing = True
        app.lasso_points = [(2.0, 2.0), (8.0, 2.0)]
        app.lasso_region_index = None
        app.regions = []
        app.selected_index = None
        app.selected_vertex = None
        app.hovered_vertex = None
        app.dragging_vertex = None
        app.region_list = Mock()
        app.status_var = Mock()
        app._sync_region_controls = Mock()
        app._update_statistics = Mock()
        app._refresh_regions = Mock()
        app._set_canvas_cursor = Mock()
        app._render = Mock()

        result = app._on_escape()

        self.assertEqual(result, "break")
        self.assertFalse(app.drawing)
        self.assertEqual(app.lasso_points, [])
        self.assertIsNone(app.lasso_region_index)
        self.assertIsNone(app.selected_index)
        self.assertIsNone(app.selected_vertex)
        app._render.assert_called()

    def test_escape_clears_active_lasso_preview_without_restoring_yellow_points(self) -> None:
        from app import ImageAnalyzerApp

        app = ImageAnalyzerApp.__new__(ImageAnalyzerApp)
        region = Region(
            "Fourth",
            [(2.0, 2.0), (18.0, 2.0), (18.0, 18.0), (2.0, 18.0)],
            "#ff0000",
        )
        app.delete_vertex_mode = False
        app.drawing = True
        app.lasso_points = region.points
        app.lasso_region_index = 0
        app.regions = [region]
        app.selected_index = 0
        app.selected_vertex = None
        app.hovered_vertex = None
        app.dragging_vertex = None
        app.region_list = Mock()
        app.status_var = Mock()
        app._sync_region_controls = Mock()
        app._sync_opacity_control = Mock()
        app._sync_border_width_control = Mock()
        app._update_statistics = Mock()
        app._refresh_regions = Mock()
        app._set_canvas_cursor = Mock()
        app._render = Mock()
        app.canvas = Mock()
        app.image = Image.new("RGB", (24, 24), "white")
        app.image_x = 0.0
        app.image_y = 0.0
        app.scale = 1.0
        app.scale_x = 1.0
        app.scale_y = 1.0
        app.suppress_vertex_hover = False

        app._on_escape()
        app._on_release(SimpleNamespace(x=2.0, y=2.0))
        app._on_mouse_move(SimpleNamespace(x=2.0, y=2.0))

        self.assertFalse(app.drawing)
        self.assertEqual(app.lasso_points, [])
        self.assertIsNone(app.lasso_region_index)
        self.assertIsNone(app.selected_index)
        self.assertIsNone(app.selected_vertex)
        self.assertIsNone(app.hovered_vertex)
        self.assertTrue(app.suppress_vertex_hover)



if __name__ == "__main__":
    unittest.main()
